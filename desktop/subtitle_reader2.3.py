import sys
import ctypes
import tkinter as tk
from tkinter import filedialog, colorchooser, font, ttk, messagebox, simpledialog
import re
import os
import json


try:
    from docx import Document
    from docx.shared import Pt
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except ImportError:
    HAS_DND = False


LANG = {
    "zh": {
        "app_title": "字幕阅读器",
        "menu_file": "文件",
        "menu_open": "打开...",
        "menu_export": "导出...",
        "menu_exit": "退出",
        "menu_lang": "语言",
        "lang_zh": "中文",
        "lang_en": "English",
        "open_file": "打开文件",
        "timecode": "时间码",
        "line_no": "行号",
        "duration": "时长",
        "font": "字体：",
        "size": "字号：",
        "text_color": "文字颜色",
        "bg_color": "背景颜色",
        "margin": "边距：",
        "status_no_file": "未加载文件",
        "status_items": "{} 条字幕",
        "status_duration": "总时长：{}",
        "no_file": "未加载字幕。\n打开或将 .srt/.sub/.ass 文件拖入窗口开始阅读。",
        "export_title": "导出字幕内容",
        "export_ok": "已导出到：\n{}",
        "export_err": "导出错误",
        "export_none": "未加载字幕文件。",
        "text_color_title": "文字颜色",
        "bg_color_title": "背景颜色",
        "sub_count": "{} 条字幕",
        "open_err": "打开文件出错：{}",
        "file_filter": "字幕文件",
        "all_files": "所有文件",
        "filetype_txt": "文本文件",
        "filetype_docx": "Word 文档",
        "export_fmt": "导出格式",
        "export_txt": "文本 (.txt)",
        "export_docx": "Word (.docx)",
        "menu_close": "关闭",
        "menu_recent": "最近打开",
        "file_missing": "(文件不存在)",
        "clear_recent": "清除记录",
        "clear_recent_confirm": "确定清除所有最近打开记录？",
        "jump_to": "跳转到位置",
        "jump_to_hint": "请输入百分比 (0-100)：",
        "menu_display": "显示模式",
        "mode_percent": "百分比%",
        "mode_page": "分页",
        "page_fmt": "第{}页/共{}页",
        "page_jump_hint": "请输入页数 (1-{})：",
        "menu_edit": "编辑字幕",
        "edit_mode": "编辑",
        "edit_title": "字幕编辑器",
        "edit_no_file": "请先打开一个字幕文件。",
        "edit_no_select": "请先选择要修改的字幕。",
        "edit_add": "新增",
        "edit_del": "删除",
        "edit_up": "上移",
        "edit_down": "下移",
        "edit_save": "保存",
        "edit_save_as": "另存为...",
        "edit_find_replace": "查找/替换",
        "edit_find_title": "查找 / 替换",
        "edit_find_what": "查找内容：",
        "edit_replace_with": "替换为：",
        "edit_find_time": "包含开始/结束时间",
        "edit_find_next": "查找下一处",
        "edit_find_prev": "查找上一处",
        "edit_find_replace_one": "替换当前",
        "edit_find_replace_all": "全部替换",
        "edit_find_not_found": "未找到：{}",
        "edit_find_found": "找到 {} 处",
        "edit_find_no_text": "请输入要查找的内容。",
        "edit_find_replaced": "已替换 {} 处。",
        "edit_find_case": "区分大小写",
        "edit_apply": "应用修改",
        "edit_start": "开始时间",
        "edit_end": "结束时间",
        "edit_text": "字幕文本",
        "col_index": "行号",
        "col_start": "开始",
        "col_end": "结束",
        "col_text": "文本",
        "edit_invalid_time": "时间格式无效：{}\n应为 HH:MM:SS,mmm",
        "edit_confirm": "确认",
        "edit_del_confirm": "确定删除选中的 {} 条字幕？",
        "edit_overwrite": "将修改写回原文件？\n{}",
        "edit_saved": "已保存到：\n{}",
        "edit_save_err": "保存失败：{}",
        "edit_new_text": "（新字幕）",
        "ed_size_list": "列表字号：",
        "ed_size_edit": "文本字号：",
        "ed_theme": "隔行：",
        "ed_odd_bg": "奇行底",
        "ed_even_bg": "偶行底",
        "ed_odd_fg": "奇字",
        "ed_even_fg": "偶字",
        "ed_color_title": "选择颜色",
        "menu_help": "帮助",
        "menu_about": "关于",
        "about_title": "关于字幕阅读器",
        "about_info": "字幕阅读器 v2.3\nSubtitle-Reader v2.3\n\n作者：武东东\n开源地址\nhttps://github.com/wdd500/subtitle-reader",
    },
    "en": {
        "app_title": "Subtitle Reader",
        "menu_file": "File",
        "menu_open": "Open...",
        "menu_export": "Export...",
        "menu_exit": "Exit",
        "menu_lang": "Language",
        "lang_zh": "中文",
        "lang_en": "English",
        "open_file": "Open File",
        "timecode": "Timecode",
        "line_no": "Line No.",
        "duration": "Duration",
        "font": "Font:",
        "size": "Size:",
        "text_color": "Text Color",
        "bg_color": "Bg Color",
        "margin": "Margin:",
        "status_no_file": "No file loaded",
        "status_items": "{} subtitles",
        "status_duration": "Total: {}",
        "no_file": "No subtitles loaded.\nOpen or drop a .srt / .sub / .ass file to begin.",
        "export_title": "Export Subtitle Content",
        "export_ok": "Exported to:\n{}",
        "export_err": "Export Error",
        "export_none": "No subtitle file loaded.",
        "text_color_title": "Text Color",
        "bg_color_title": "Background Color",
        "sub_count": "{} subtitles",
        "open_err": "Error opening file: {}",
        "file_filter": "Subtitle files",
        "all_files": "All files",
        "filetype_txt": "Text file",
        "filetype_docx": "Word document",
        "export_fmt": "Export Format",
        "export_txt": "Text (.txt)",
        "export_docx": "Word (.docx)",
        "menu_close": "Close",
        "menu_recent": "Recent Files",
        "file_missing": "(file missing)",
        "clear_recent": "Clear Recent Files",
        "clear_recent_confirm": "Clear all recent files?",
        "jump_to": "Jump to Position",
        "jump_to_hint": "Enter percentage (0-100):",
        "menu_display": "Display Mode",
        "mode_percent": "Percentage %",
        "mode_page": "Paging",
        "page_fmt": "Page {}/{}",
        "page_jump_hint": "Enter page (1-{}):",
        "menu_edit": "Edit Subtitles",
        "edit_mode": "Edit",
        "edit_title": "Subtitle Editor",
        "edit_no_file": "Open a subtitle file first.",
        "edit_no_select": "Select a subtitle to edit first.",
        "edit_add": "Add",
        "edit_del": "Delete",
        "edit_up": "Move Up",
        "edit_down": "Move Down",
        "edit_save": "Save",
        "edit_save_as": "Save As...",
        "edit_find_replace": "Find/Replace",
        "edit_find_title": "Find / Replace",
        "edit_find_what": "Find what:",
        "edit_replace_with": "Replace with:",
        "edit_find_time": "Include start/end time",
        "edit_find_next": "Find Next",
        "edit_find_prev": "Find Prev",
        "edit_find_replace_one": "Replace",
        "edit_find_replace_all": "Replace All",
        "edit_find_not_found": "Not found: {}",
        "edit_find_found": "{} found",
        "edit_find_no_text": "Enter text to find.",
        "edit_find_replaced": "Replaced {} occurrence(s).",
        "edit_find_case": "Match case",
        "edit_apply": "Apply",
        "edit_start": "Start Time",
        "edit_end": "End Time",
        "edit_text": "Subtitle Text",
        "col_index": "#",
        "col_start": "Start",
        "col_end": "End",
        "col_text": "Text",
        "edit_invalid_time": "Invalid time: {}\nExpected HH:MM:SS,mmm",
        "edit_confirm": "Confirm",
        "edit_del_confirm": "Delete {} selected subtitle(s)?",
        "edit_overwrite": "Overwrite original file?\n{}",
        "edit_saved": "Saved to:\n{}",
        "edit_save_err": "Save failed: {}",
        "edit_new_text": "(new subtitle)",
        "ed_size_list": "List size:",
        "ed_size_edit": "Text size:",
        "ed_theme": "Stripe:",
        "ed_odd_bg": "Odd bg",
        "ed_even_bg": "Even bg",
        "ed_odd_fg": "Odd fg",
        "ed_even_fg": "Even fg",
        "ed_color_title": "Pick color",
        "menu_help": "Help",
        "menu_about": "About",
        "about_title": "About Subtitle Reader",
        "about_info": "字幕阅读器 v2.3\nSubtitle-Reader v2.3\n\n作者：武东东\n开源地址\nhttps://github.com/wdd500/subtitle-reader",
    },
}


def enable_high_dpi():
    if sys.platform != "win32":
        return
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


def get_dpi_scale():
    try:
        hwnd = ctypes.windll.user32.GetDesktopWindow()
        dpi = ctypes.windll.user32.GetDpiForWindow(hwnd)
        return dpi / 96.0
    except Exception:
        return 1.0


def get_work_area():
    class RECT(ctypes.Structure):
        _fields_ = [("left", ctypes.c_long), ("top", ctypes.c_long),
                    ("right", ctypes.c_long), ("bottom", ctypes.c_long)]
    rect = RECT()
    try:
        ok = ctypes.windll.user32.SystemParametersInfoW(
            0x0030, 0, ctypes.byref(rect), 0)
        if ok:
            return {
                "x": rect.left,
                "y": rect.top,
                "width": rect.right - rect.left,
                "height": rect.bottom - rect.top,
            }
    except Exception:
        pass
    return None


def _window_outer(root):
    class RECT(ctypes.Structure):
        _fields_ = [("left", ctypes.c_long), ("top", ctypes.c_long),
                    ("right", ctypes.c_long), ("bottom", ctypes.c_long)]
    try:
        hwnd = root.winfo_id()
        rect = RECT()
        if ctypes.windll.user32.GetWindowRect(hwnd, ctypes.byref(rect)):
            return rect.left, rect.top, rect.right - rect.left, rect.bottom - rect.top
    except Exception:
        pass
    return None


SUBTITLE_EXTS = ('.srt', '.sub', '.ass', '.ssa')


class SubtitleItem:
    def __init__(self, index=0, start="", end="", text="", duration=""):
        self.index = index
        self.start = start
        self.end = end
        self.text = text
        self.duration = duration


class SubtitleParser:
    @staticmethod
    def parse_srt(content):
        items = []
        blocks = re.split(r'\n\s*\n', content.strip())
        for block in blocks:
            lines = block.strip().split('\n')
            if len(lines) >= 3:
                try:
                    idx = int(lines[0])
                    time_match = re.match(
                        r'(\d+:\d+:\d+[,.]\d+)\s*-->\s*(\d+:\d+:\d+[,.]\d+)', lines[1])
                    if time_match:
                        start = time_match.group(1).replace(',', '.')
                        end = time_match.group(2).replace(',', '.')
                        text = '\n'.join(lines[2:])
                        duration = SubtitleParser._calc_duration(start, end)
                        items.append(SubtitleItem(idx, start, end, text, duration))
                except ValueError:
                    continue
        return items

    @staticmethod
    def parse_sub(content):
        items = []
        lines = content.strip().split('\n')
        for i, line in enumerate(lines):
            line = line.strip()
            if not line:
                continue
            match = re.match(r'\{(\d+)\}\{(\d+)\}(.*)', line)
            if match:
                start_frame = match.group(1)
                end_frame = match.group(2)
                text = match.group(3).replace('|', '\n')
                duration = str(int(end_frame) - int(start_frame)) + ' frames'
                items.append(SubtitleItem(i + 1, start_frame, end_frame, text, duration))
        return items

    @staticmethod
    def parse_ass(content):
        items = []
        events_section = False
        format_line = None
        for line in content.split('\n'):
            line = line.strip()
            if line.startswith('[Events]'):
                events_section = True
                continue
            if events_section:
                if line.startswith('Format:'):
                    format_line = line[7:].strip().split(',')
                    format_line = [f.strip() for f in format_line]
                    continue
                if line.startswith('Dialogue:'):
                    parts = line[9:].strip()
                    if format_line:
                        vals = SubtitleParser._split_ass(parts)
                        data = {}
                        for j, key in enumerate(format_line):
                            if j < len(vals):
                                data[key] = vals[j]
                        start = data.get('Start', '')
                        end = data.get('End', '')
                        text = data.get('Text', '')
                        text = re.sub(r'\{[^}]*\}', '', text)
                        text = text.replace('\\N', '\n').replace('\\n', '\n')
                        duration = SubtitleParser._calc_duration_ass(start, end)
                        items.append(SubtitleItem(len(items) + 1, start, end, text, duration))
        return items

    @staticmethod
    def _split_ass(line):
        parts = []
        current = ''
        in_braces = 0
        for ch in line:
            if ch == '{':
                in_braces += 1
                current += ch
            elif ch == '}':
                in_braces -= 1
                current += ch
            elif ch == ',' and in_braces == 0:
                parts.append(current.strip())
                current = ''
            else:
                current += ch
        parts.append(current.strip())
        return parts

    @staticmethod
    def _calc_duration(start, end):
        def to_seconds(t):
            parts = t.replace(',', '.').split(':')
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
        try:
            d = to_seconds(end) - to_seconds(start)
            return f'{d:.2f}s'
        except:
            return ''

    @staticmethod
    def _calc_duration_ass(start, end):
        def to_seconds(t):
            parts = t.replace(',', '.').split(':')
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
        try:
            d = to_seconds(end) - to_seconds(start)
            return f'{d:.2f}s'
        except:
            return ''

    @staticmethod
    def parse_file(filepath):
        ext = os.path.splitext(filepath)[1].lower()
        with open(filepath, encoding='utf-8') as f:
            content = f.read()
        if ext == '.srt':
            return SubtitleParser.parse_srt(content)
        elif ext == '.sub':
            return SubtitleParser.parse_sub(content)
        elif ext in ('.ass', '.ssa'):
            return SubtitleParser.parse_ass(content)
        else:
            return []

    @staticmethod
    def to_srt_string(items):
        blocks = []
        for i, it in enumerate(items, 1):
            start = it.start.replace('.', ',')
            end = it.end.replace('.', ',')
            blocks.append(f"{i}\n{start} --> {end}\n{it.text}")
        return '\n\n'.join(blocks) + '\n'

    @staticmethod
    def to_sub_string(items):
        lines = []
        for it in items:
            lines.append("{%s}{%s}%s" % (it.start, it.end, it.text.replace('\n', '|')))
        return '\n'.join(lines) + '\n'

    @staticmethod
    def to_ass_string(items):
        head = (
            "[Script Info]\n"
            "; Generated by Subtitle-Reader\n"
            "ScriptType: v4.00+\n"
            "WrapStyle: 0\n"
            "ScaledBorderAndShadow: yes\n"
            "YCbCr Matrix: TV.709\n"
            "\n"
            "[V4+ Styles]\n"
            "Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding\n"
            "Style: Default,Microsoft YaHei,32,&H00FFFFFF,&H000000FF,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,2,0,2,10,10,10,1\n"
            "\n"
            "[Events]\n"
            "Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\n"
        )
        events = []
        for it in items:
            text = it.text.replace('\n', '\\N')
            events.append(f"Dialogue: 0,{it.start},{it.end},Default,,0,0,0,,{text}")
        return head + '\n'.join(events) + '\n'

    @staticmethod
    def write_file(filepath, items):
        ext = os.path.splitext(filepath)[1].lower()
        if ext == '.srt':
            content = SubtitleParser.to_srt_string(items)
        elif ext == '.sub':
            content = SubtitleParser.to_sub_string(items)
        else:
            content = SubtitleParser.to_ass_string(items)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)


class EdGrid(tk.Canvas):
    def __init__(self, master, app):
        tk.Canvas.__init__(self, master, bg="#FFFFFF", highlightthickness=0,
                           borderwidth=0)
        self._app = app
        self._sel = []
        self._anchor = None
        self._off = 0
        self._hx = 0
        self._sb = None
        self._hsb = None
        self._first = 0.0
        self._last = 1.0
        self._hfirst = 0.0
        self._hlast = 1.0
        self._ver = 0
        self._fk = None
        self._auto = (40, 150, 150, 400)
        self._fw = (40, 150, 150, 400)
        self._fonts = {}
        self._cw = {}
        self.bind("<Configure>", lambda e: self.redraw())
        self.bind("<Button-1>", self._press)
        self.bind("<B1-Motion>", self._drag)
        self.bind("<MouseWheel>", self._wheel)

    def attach_sb(self, sb):
        self._sb = sb

    def attach_hsb(self, hsb):
        self._hsb = hsb

    def invalidate(self):
        self._ver += 1
        self.redraw()

    def selection(self):
        return list(self._sel)

    def selection_set(self, *iids):
        sel = []
        for a in iids:
            if isinstance(a, (list, tuple)):
                sel.extend(a)
            else:
                sel.append(a)
        self._sel = [str(s) for s in sel]
        self.redraw()

    def _rowh(self):
        return max(24, int(self._app.editor_font_size * 1.7))

    def _headh(self):
        return max(28, int(self._app.ui_font_size * 1.5))

    def _maxoff(self):
        total = len(self._app.items) * self._rowh()
        view = max(0, int(self.winfo_height()) - self._headh())
        return max(0, total - view)

    def _maxhx(self):
        total = sum(self._fw)
        return max(0, total - max(1, int(self.winfo_width())))

    def _clamp(self):
        self._off = max(0, min(self._maxoff(), self._off))
        self._hx = max(0, min(self._maxhx(), self._hx))
        n = len(self._app.items)
        self._sel = [s for s in self._sel
                     if str(s).isdigit() and 0 < int(s) <= n]

    def yview(self, *args):
        if not args:
            return self._first, self._last
        try:
            if args[0] == "moveto":
                self._off = int(float(args[1]) * max(1, self._maxoff()))
            elif args[0] == "scroll":
                amt = int(args[1])
                if args[2] == "units":
                    self._off += amt * self._rowh()
                else:
                    view = max(1, int(self.winfo_height()))
                    self._off += amt * view
        except Exception:
            pass
        self.redraw()

    def hview(self, *args):
        if not args:
            return self._hfirst, self._hlast
        try:
            if args[0] == "moveto":
                self._hx = int(float(args[1]) * max(1, self._maxhx()))
            elif args[0] == "scroll":
                amt = int(args[1])
                if args[2] == "units":
                    self._hx += amt * 60
                else:
                    view = max(1, int(self.winfo_width()))
                    self._hx += amt * view
        except Exception:
            pass
        self.redraw()

    def _wheel(self, ev):
        if ev.state & 1:
            step = int(ev.delta / max(1, abs(ev.delta))) * 180
            self._hx -= step
            self.redraw()
            return
        step = int(ev.delta / max(1, abs(ev.delta))) * self._rowh() * 3
        self._off -= step
        self.redraw()

    def _row_at(self, y):
        hh = self._headh()
        if y < hh:
            return None
        cy = y - hh + self._off
        if cy < 0:
            return None
        return int(cy // self._rowh())

    def _font_for(self, family, size):
        key = (family, size)
        f = self._fonts.get(key)
        if f is None:
            f = font.Font(root=self, family=family, size=size)
            self._fonts[key] = f
        return f

    def _char_w(self, key, ch):
        cw = self._cw.setdefault(key, {})
        w = cw.get(ch)
        if w is None:
            f = self._fonts.get(key)
            if f is None:
                f = self._font_for(*key)
            try:
                w = int(f.measure(ch))
            except Exception:
                w = 0
            cw[ch] = w
        return w

    def _str_w(self, key, s):
        cw = self._cw.setdefault(key, {})
        f = self._fonts.get(key)
        if f is None:
            f = self._font_for(*key)
        total = 0
        for ch in s:
            w = cw.get(ch)
            if w is None:
                try:
                    w = int(f.measure(ch))
                except Exception:
                    w = 0
                cw[ch] = w
            total += w
        return total

    def _auto_widths(self, hkey, rkey, hf):
        a = self._app
        labels = (a._("col_index"), a._("col_start"), a._("col_end"), a._("col_text"))
        req = [max(40, int(hf.measure(labels[i])) + 12) for i in range(4)]
        caps = [220, 420, 420, 12000]
        sw = self._str_w
        for i, it in enumerate(a.items, 1):
            w0 = sw(rkey, str(i)) + 12
            w1 = sw(rkey, it.start) + 12
            w2 = sw(rkey, it.end) + 12
            w3 = sw(rkey, it.text.replace("\n", " ")) + 16
            if w0 > req[0]:
                req[0] = w0
            if w1 > req[1]:
                req[1] = w1
            if w2 > req[2]:
                req[2] = w2
            if w3 > req[3]:
                req[3] = w3
        return [min(req[i], caps[i]) for i in range(4)]

    def _press(self, ev):
        self.focus_set()
        idx = self._row_at(ev.y)
        if idx is None or idx >= len(self._app.items):
            self._sel = []
            self._anchor = None
            self.redraw()
            self._notify()
            return
        iid = str(idx + 1)
        if ev.state & 4:
            if iid in self._sel:
                self._sel = [s for s in self._sel if s != iid]
            else:
                self._sel = self._sel + [iid]
            self._anchor = idx
        elif ev.state & 1:
            start = self._anchor if self._anchor is not None else idx
            lo, hi = min(start, idx), max(start, idx)
            self._sel = [str(i + 1) for i in range(lo, hi + 1)]
        else:
            self._sel = [iid]
            self._anchor = idx
        self.redraw()
        self._notify()

    def _drag(self, ev):
        if self._anchor is None:
            return
        idx = self._row_at(ev.y)
        n = len(self._app.items)
        if idx is None:
            idx = n - 1
        if idx > n - 1:
            idx = n - 1
        if idx < 0:
            return
        lo, hi = min(self._anchor, idx), max(self._anchor, idx)
        self._sel = [str(i + 1) for i in range(lo, hi + 1)]
        self.redraw()

    def _notify(self):
        try:
            self._app._ed_on_select()
        except Exception:
            pass

    def redraw(self):
        a = self._app
        W = int(self.winfo_width())
        H = int(self.winfo_height())
        if W <= 1 or H <= 1:
            return
        labels = (a._("col_index"), a._("col_start"), a._("col_end"), a._("col_text"))
        hkey = (a.ui_font_family, a.ui_font_size)
        rkey = (a.font_family, a.editor_font_size)
        hf = self._font_for(*hkey)
        rf = self._font_for(*rkey)
        fkey = (hkey, rkey, self._ver)
        if fkey != self._fk:
            self._fk = fkey
            self._auto = self._auto_widths(hkey, rkey, hf)
        caps = [220, 420, 420]
        fixed = [min(self._auto[i], caps[i]) for i in range(3)]
        text = self._auto[3]
        fill = W - sum(fixed)
        if fill > text:
            text = fill
        fw = fixed + [text]
        self._fw = tuple(fw)
        self._clamp()
        off = self._off
        hx = self._hx
        hh = self._headh()
        rh = self._rowh()
        self.delete("all")
        xs = [0]
        for w in fw:
            xs.append(xs[-1] + w)
        self.create_rectangle(0, 0, W - 1, H - 1, fill="#FFFFFF", outline="")
        self.create_rectangle(0, 0, W - 1, hh - 1, fill="#F0F0F0", outline="")
        hy = (hh - 1) // 2
        for ci in range(4):
            self.create_text(xs[ci] + 4 - hx, hy, text=labels[ci], anchor=tk.W,
                             font=hf, fill="#000000")
        items = a.items
        cy0 = hh - off
        idx = max(0, (off - hh) // rh)
        while idx < len(items):
            y0 = cy0 + idx * rh
            if y0 >= H:
                break
            y1 = y0 + rh
            if y1 < 0:
                idx += 1
                continue
            it = items[idx]
            odd = (idx % 2 == 0)
            bg = a.ed_odd_bg if odd else a.ed_even_bg
            fg = a.ed_odd_fg if odd else a.ed_even_fg
            if str(idx + 1) in self._sel:
                fillc, text_fill = fg, bg
            else:
                fillc, text_fill = bg, fg
            self.create_rectangle(0, y0, W - 1, y1 - 1, fill=fillc, outline="")
            ymid = (y0 + y1 - 1) // 2
            vals = (str(idx + 1), it.start, it.end)
            for ci in range(3):
                self.create_text(xs[ci] + 4 - hx, ymid, text=vals[ci],
                                 anchor=tk.W, font=rf, fill=text_fill)
            self.create_text(xs[3] + 4 - hx, ymid,
                             text=it.text.replace("\n", " "), anchor=tk.W,
                             font=rf, fill=text_fill)
            idx += 1
        for ci in range(5):
            self.create_line(xs[ci] - hx, 0, xs[ci] - hx, H - 1, fill="#000000")
        b = 0
        while True:
            y = cy0 + b * rh
            if y >= H:
                break
            if y > 0:
                self.create_line(0, y, W - 1, y, fill="#000000")
            b += 1
        self.create_line(0, H - 1, W - 1, H - 1, fill="#000000")
        content = len(items) * rh
        view = max(0, H - hh)
        span = max(content, view)
        self._first = off / max(1, span)
        self._last = min(1.0, (off + H) / max(1, span))
        total_x = sum(fw)
        span_x = max(total_x, W)
        self._hfirst = hx / max(1, span_x)
        self._hlast = min(1.0, (hx + W) / max(1, span_x))
        if self._sb is not None:
            try:
                self._sb.set(self._first, self._last)
            except Exception:
                pass
        if self._hsb is not None:
            try:
                self._hsb.set(self._hfirst, self._hlast)
            except Exception:
                pass


class SubtitleReader:
    _TIME_RE = re.compile(r'^\d{1,3}:\d{2}:\d{2}[,.]\d{1,3}$')

    def __init__(self, root):
        self.root = root
        self.current_lang = "zh"
        self.items = []
        self.current_filepath = None
        self.show_time = False
        self.show_index = False
        self.show_duration = False
        self.text_color = "#000000"
        self.bg_color = "#FFFFFF"
        self.font_size = 22
        self.font_family = "Segoe UI"
        self.ui_font_family = "梦源黑体 SC W15"
        self.ui_font_size = 30
        self.margin = 10
        self.editor_font_size = 30
        self.editor_text_size = 50
        self.ed_odd_bg = "#FFFFFF"
        self.ed_even_bg = "#F2F2F2"
        self.ed_odd_fg = "#000000"
        self.ed_even_fg = "#000000"
        self.w = {}
        self._last_rebuild_recent = []
        self.display_mode = tk.StringVar(value="percent")
        self.build_ui()
        self.apply_language()
        self.load_settings()
        if HAS_DND:
            self.register_drop_targets(self.root)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---- i18n helpers ----
    def _(self, key):
        return LANG[self.current_lang].get(key, key)

    def apply_language(self):
        self.root.title(self._("app_title"))
        self.rebuild_file_menu()

        self.lang_menu.delete(0, "end")
        self.lang_menu.add_command(label=self._("lang_zh"), command=lambda: self.set_language("zh"))
        self.lang_menu.add_command(label=self._("lang_en"), command=lambda: self.set_language("en"))

        self.display_menu.delete(0, "end")
        self.display_menu.add_radiobutton(label=self._("mode_percent"), variable=self.display_mode, value="percent", command=self.on_display_mode_change)
        self.display_menu.add_radiobutton(label=self._("mode_page"), variable=self.display_mode, value="page", command=self.on_display_mode_change)

        self.help_menu.delete(0, "end")
        self.help_menu.add_command(label=self._("menu_about"), command=self.show_about)

        self.menubar.delete(0, "end")
        self.menubar.add_cascade(label=self._("menu_file"), menu=self.file_menu)
        self.menubar.add_cascade(label=self._("menu_display"), menu=self.display_menu)
        self.menubar.add_cascade(label=self._("menu_lang"), menu=self.lang_menu)
        self.menubar.add_cascade(label=self._("menu_help"), menu=self.help_menu)
        self.root.config(menu=self.menubar)

        self.w["open_btn"].config(text=self._("open_file"))
        self.w["edit_btn"].config(text=self._("edit_mode"))
        self.w["time_cb"].config(text=self._("timecode"))
        self.w["idx_cb"].config(text=self._("line_no"))
        self.w["dur_cb"].config(text=self._("duration"))
        self.w["font_lbl"].config(text=self._("font"))
        self.w["size_lbl"].config(text=self._("size"))
        self.w["txt_color_btn"].config(text=self._("text_color"))
        self.w["bg_color_btn"].config(text=self._("bg_color"))
        self.w["margin_lbl"].config(text=self._("margin"))

        self.refresh_display()
        self.refresh_status_bar()

    # ---- build UI ----
    def build_ui(self):
        self.menubar = tk.Menu(self.root)
        self.file_menu = tk.Menu(self.menubar, tearoff=0)
        self.file_menu.add_command(label="", command=self.open_file, accelerator="Ctrl+O")
        self.file_menu.add_separator()
        self.file_menu.add_command(label="", command=self.export_content, accelerator="Ctrl+E")
        self.file_menu.add_separator()
        self.file_menu.add_command(label="", command=self.root.quit)
        self.menubar.add_cascade(label="", menu=self.file_menu)

        self.display_menu = tk.Menu(self.menubar, tearoff=0)
        self.display_menu.add_radiobutton(label="", variable=self.display_mode, value="percent", command=self.on_display_mode_change)
        self.display_menu.add_radiobutton(label="", variable=self.display_mode, value="page", command=self.on_display_mode_change)
        self.menubar.add_cascade(label="", menu=self.display_menu)

        self.lang_menu = tk.Menu(self.menubar, tearoff=0)
        self.lang_menu.add_command(label="", command=lambda: self.set_language("zh"))
        self.lang_menu.add_command(label="", command=lambda: self.set_language("en"))
        self.menubar.add_cascade(label="", menu=self.lang_menu)

        self.help_menu = tk.Menu(self.menubar, tearoff=0)
        self.help_menu.add_command(label="", command=self.show_about)
        self.menubar.add_cascade(label="", menu=self.help_menu)

        self.root.config(menu=self.menubar)
        self.set_window_icon()
        self.root.bind_all("<Control-o>", lambda e: self.open_file())
        self.root.bind_all("<Control-e>", lambda e: self.export_content())

        toolbar = ttk.Frame(self.root)
        toolbar.pack(fill=tk.X, padx=3, pady=3)
        self.toolbar = toolbar
        try:
            tool_style = ttk.Style(self.root)
            tool_style.configure("UIFont22.TCombobox", font=(self.ui_font_family, 22))
            tool_style.configure("UIFont22.TSpinbox", font=(self.ui_font_family, 22))
            tool_style.configure("UIFont40.TSpinbox", font=(self.ui_font_family, 40))
        except Exception:
            pass

        self.w["open_btn"] = ttk.Button(toolbar, command=self.open_file)
        self.w["open_btn"].pack(side=tk.LEFT, padx=2)

        self.w["edit_btn"] = ttk.Button(toolbar, command=self.open_editor)
        self.w["edit_btn"].pack(side=tk.LEFT, padx=2)

        self.time_var = tk.BooleanVar()
        self.w["time_cb"] = ttk.Checkbutton(toolbar, variable=self.time_var, command=self.refresh_display)
        self.w["time_cb"].pack(side=tk.LEFT, padx=2)

        self.index_var = tk.BooleanVar()
        self.w["idx_cb"] = ttk.Checkbutton(toolbar, variable=self.index_var, command=self.refresh_display)
        self.w["idx_cb"].pack(side=tk.LEFT, padx=2)

        self.dur_var = tk.BooleanVar()
        self.w["dur_cb"] = ttk.Checkbutton(toolbar, variable=self.dur_var, command=self.refresh_display)
        self.w["dur_cb"].pack(side=tk.LEFT, padx=2)

        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)

        self.w["font_lbl"] = ttk.Label(toolbar)
        self.w["font_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        fonts = sorted(font.families())
        self.font_combo = ttk.Combobox(toolbar, values=fonts, width=32, style="UIFont22.TCombobox")
        self.font_combo.configure(font=(self.ui_font_family, 22))
        self.font_combo.set(self.font_family)
        self.font_combo.pack(side=tk.LEFT, padx=2)
        self.font_combo.bind("<<ComboboxSelected>>", self.apply_font)

        self.w["size_lbl"] = ttk.Label(toolbar)
        self.w["size_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        self.size_var = tk.StringVar(value=str(self.font_size))
        self.size_spin = ttk.Spinbox(toolbar, from_=8, to=72, width=8,
                              textvariable=self.size_var, style="UIFont22.TSpinbox")
        self.size_spin.configure(font=(self.ui_font_family, 22))
        self.size_spin.pack(side=tk.LEFT, padx=2)
        self.size_var.trace_add("write", self.apply_font)

        toolbar2 = ttk.Frame(self.root)
        toolbar2.pack(fill=tk.X, padx=3, pady=(0, 3))

        self.w["txt_color_btn"] = ttk.Button(toolbar2, command=self.choose_text_color)
        self.w["txt_color_btn"].pack(side=tk.LEFT, padx=2)

        self.w["bg_color_btn"] = ttk.Button(toolbar2, command=self.choose_bg_color)
        self.w["bg_color_btn"].pack(side=tk.LEFT, padx=2)

        ttk.Separator(toolbar2, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)

        self.w["margin_lbl"] = ttk.Label(toolbar2)
        self.w["margin_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        self.margin_var = tk.StringVar(value=str(self.margin))
        self.margin_spin = ttk.Spinbox(toolbar2, from_=0, to=200, width=8,
                                       textvariable=self.margin_var, style="UIFont22.TSpinbox",
                                       command=self.apply_margin)
        self.margin_spin.configure(font=(self.ui_font_family, 22))
        self.margin_spin.pack(side=tk.LEFT, padx=2)
        self.margin_var.trace_add("write", self.apply_margin)

        prog_frame = ttk.Frame(self.root)
        prog_frame.pack(fill=tk.X, padx=5, pady=(0, 2))
        self.progress_bar = ttk.Progressbar(prog_frame, mode="determinate", value=0)
        self.progress_bar.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.progress_bar.bind("<Button-1>", self._on_progress_click)
        self.progress_label = ttk.Label(prog_frame, text="0%", anchor=tk.CENTER,
                                        font=(self.ui_font_family, 10, "bold"))
        self.progress_label.pack(side=tk.RIGHT, padx=(5, 0))
        self.progress_label.bind("<Button-1>", self._on_progress_label_click)

        status_frame = ttk.Frame(self.root, padding=(8, 3))
        status_frame.pack(fill=tk.X, side=tk.BOTTOM)
        self.status_bar = ttk.Label(status_frame, text="", anchor=tk.W)
        self.status_bar.pack(side=tk.LEFT, fill=tk.X, expand=True)

        text_frame = ttk.Frame(self.root, padding=5)
        text_frame.pack(fill=tk.BOTH, expand=True)

        self.text_area = tk.Text(text_frame, wrap=tk.WORD, font=(self.font_family, self.font_size),
                                 fg=self.text_color, bg=self.bg_color, relief=tk.SUNKEN, borderwidth=2)
        self.scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self._on_scrollbar)
        self.text_area.configure(yscrollcommand=self._on_text_scroll)
        self.text_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.text_area.config(state=tk.DISABLED)
        self.text_area.bind("<MouseWheel>", self._on_mousewheel)
        self.text_area.bind("<ButtonRelease-1>", self.update_progress)
        self.text_area.bind("<KeyRelease>", self.update_progress)

    def set_language(self, lang):
        self.current_lang = lang
        self.apply_language()

    # ---- file ----
    def open_file(self, filepath=None, restore_pos=False):
        if filepath is None:
            path = filedialog.askopenfilename(
                title=self._("menu_open"),
                filetypes=[(self._("file_filter"), "*.srt *.sub *.ass *.ssa"), (self._("all_files"), "*.*")]
            )
        else:
            path = filepath
        if not path:
            return
        self.save_position()
        try:
            self.items = SubtitleParser.parse_file(path)
            self.current_filepath = path
            self.root.title(f"{self._('app_title')} - {os.path.basename(path)}")
            self.refresh_display()
            self.refresh_status_bar()
            win = getattr(self, "_edit_win", None)
            if win is not None and win.winfo_exists():
                self._ed_refresh_tree()
            if restore_pos:
                self.restore_position()
            self.update_recent_files(path)
            self.update_progress()
            self.rebuild_file_menu()
        except Exception as e:
            self.text_area.config(state=tk.NORMAL)
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, self._("open_err").format(e))
            self.text_area.config(state=tk.DISABLED)
            self.refresh_status_bar()
            self.rebuild_file_menu()

    def on_drop_files(self, paths):
        target = next((p for p in paths
                       if os.path.splitext(p)[1].lower() in SUBTITLE_EXTS), None)
        if target is None and paths:
            target = paths[0]
        if target:
            self.open_file(filepath=target)

    def register_drop_targets(self, widget):
        if isinstance(widget, tk.Menu):
            return
        try:
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", self.on_dnd_drop)
        except Exception:
            return
        for child in widget.winfo_children():
            self.register_drop_targets(child)

    def on_dnd_drop(self, event):
        try:
            paths = self.root.tk.splitlist(event.data)
        except Exception:
            paths = event.data
        if isinstance(paths, str):
            paths = [paths]
        else:
            paths = list(paths)
        self.on_drop_files(paths)
        return "copy"

    def refresh_display(self):
        self.show_time = self.time_var.get()
        self.show_index = self.index_var.get()
        self.show_duration = self.dur_var.get()
        self.text_area.config(state=tk.NORMAL)
        self.text_area.delete(1.0, tk.END)
        if not self.items:
            self.text_area.insert(tk.END, self._("no_file"))
        else:
            for item in self.items:
                parts = []
                if self.show_index:
                    parts.append(f"[{item.index}]")
                if self.show_time:
                    parts.append(f"{item.start} -> {item.end}")
                line = item.text
                if self.show_duration and item.duration:
                    parts.append(f"(dur: {item.duration})")
                prefix = ' '.join(parts)
                if prefix:
                    self.text_area.insert(tk.END, prefix + '\n')
                self.text_area.insert(tk.END, line + '\n\n')
        self.text_area.config(state=tk.DISABLED)
        self.progress_bar.configure(maximum=100)
        self.progress_bar["value"] = 0

    def _format_duration(self, seconds):
        seconds = max(0, int(seconds))
        h, rem = divmod(seconds, 3600)
        m, s = divmod(rem, 60)
        if h:
            return f"{h}:{m:02d}:{s:02d}"
        return f"{m}:{s:02d}"

    def _total_duration(self):
        total = 0.0
        for item in self.items:
            if item.start and item.end:
                try:
                    def to_sec(t):
                        if ":" not in t:
                            return 0.0
                        parts = t.replace(",", ".").split(":")
                        return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
                    total += max(0.0, to_sec(item.end) - to_sec(item.start))
                except Exception:
                    continue
        return total

    def refresh_status_bar(self):
        if not self.items or not self.current_filepath:
            self.status_bar.config(text=self._("status_no_file"))
            return
        name = os.path.basename(self.current_filepath)
        text = f"{self._('status_items').format(len(self.items))}"
        dur = self._total_duration()
        if dur > 0:
            text += f"  |  {self._('status_duration').format(self._format_duration(dur))}"
        self.status_bar.config(text=f"{name}  -  {text}")

    def apply_font(self, *args):
        try:
            self.font_size = int(self.size_var.get())
        except:
            self.font_size = 22
        self.font_family = self.font_combo.get() or "Segoe UI"
        self.set_display_style()

    def apply_margin(self, *args):
        try:
            m = int(self.margin_var.get())
        except Exception:
            m = self.margin
        m = max(0, min(200, m))
        self.margin = m
        self.text_area.config(padx=m)

    def choose_text_color(self):
        color = colorchooser.askcolor(initialcolor=self.text_color, title=self._("text_color_title"))
        if color[1]:
            self.text_color = color[1]
            self.text_area.config(fg=self.text_color)

    def choose_bg_color(self):
        color = colorchooser.askcolor(initialcolor=self.bg_color, title=self._("bg_color_title"))
        if color[1]:
            self.bg_color = color[1]
            self.text_area.config(bg=self.bg_color)

    # ---- settings persistence ----
    def load_settings(self):
        try:
            cfg = self.load_config()
        except Exception:
            cfg = {}
        settings = cfg.get("settings", {})
        size = settings.get("font_size", self.font_size)
        try:
            size = int(size)
        except Exception:
            size = self.font_size
        family = settings.get("font_family") or self.font_family
        if family not in font.families():
            family = self.font_family
        self.font_size = size
        self.font_family = family
        self.text_color = settings.get("text_color", self.text_color)
        self.bg_color = settings.get("bg_color", self.bg_color)
        try:
            self.margin = max(0, min(200, int(settings.get("margin", self.margin))))
        except Exception:
            pass
        try:
            self.editor_font_size = max(8, min(120, int(settings.get("editor_font_size", self.editor_font_size))))
        except Exception:
            pass
        try:
            self.editor_text_size = max(8, min(120, int(settings.get("editor_text_size", self.editor_text_size))))
        except Exception:
            pass
        for k in ("ed_odd_bg", "ed_even_bg", "ed_odd_fg", "ed_even_fg"):
            v = settings.get(k)
            if isinstance(v, str) and len(v) == 7 and v[0] == "#":
                setattr(self, k, v)
        lang = settings.get("language")
        if lang in ("zh", "en"):
            self.current_lang = lang
        self.font_combo.set(family)
        self.size_var.set(str(size))
        self.margin_var.set(str(self.margin))
        self.set_display_style()
        if lang in ("zh", "en"):
            self.apply_language()

    def set_display_style(self):
        self.text_area.config(font=(self.font_family, self.font_size),
                              fg=self.text_color, bg=self.bg_color,
                              padx=self.margin)
        self.apply_ui_font()

    def apply_ui_font(self):
        fam = self.ui_font_family
        sz = self.ui_font_size
        try:
            self.menubar.config(font=(fam, sz))
            for m in (self.file_menu, self.display_menu, self.lang_menu, self.help_menu):
                m.config(font=(fam, sz))
        except Exception:
            pass
        try:
            style = ttk.Style(self.root)
            for cls in ("TLabel", "TButton", "TCheckbutton", "TSpinbox",
                        "TCombobox", "TEntry", "TProgressbar"):
                style.configure(cls, font=(fam, sz))
        except Exception:
            pass
        try:
            self.progress_label.config(font=(fam, sz, "bold"))
        except Exception:
            pass
        try:
            self.root.option_add("*TCombobox*Listbox.font", (fam, 22))
        except Exception:
            pass

    def save_settings(self):
        try:
            cfg = self.load_config()
        except Exception:
            cfg = {}
        cfg["settings"] = {
            "font_family": self.font_family,
            "font_size": self.font_size,
            "text_color": self.text_color,
            "bg_color": self.bg_color,
            "margin": self.margin,
            "editor_font_size": self.editor_font_size,
            "editor_text_size": self.editor_text_size,
            "ed_odd_bg": self.ed_odd_bg,
            "ed_even_bg": self.ed_even_bg,
            "ed_odd_fg": self.ed_odd_fg,
            "ed_even_fg": self.ed_even_fg,
            "language": self.current_lang,
        }
        self.save_config(cfg)

    def on_close(self):
        self.save_position()
        self.save_settings()
        self.root.destroy()

    # ---- export ----
    def build_export_text(self):
        lines = []
        if not self.items:
            return lines
        for item in self.items:
            parts = []
            if self.index_var.get():
                parts.append(f"[{item.index}]")
            if self.time_var.get():
                parts.append(f"{item.start} -> {item.end}")
            if self.dur_var.get() and item.duration:
                parts.append(f"(dur: {item.duration})")
            prefix = ' '.join(parts)
            if prefix:
                lines.append(prefix)
            lines.append(item.text)
            lines.append('')
        return lines

    def export_content(self):
        if not self.items:
            messagebox.showinfo(self._("export_title"), self._("export_none"))
            return
        base = os.path.splitext(os.path.basename(self.current_filepath))[0]
        path = filedialog.asksaveasfilename(
            initialfile=base,
            defaultextension=".txt",
            filetypes=[
                (self._("filetype_txt"), "*.txt"),
                (self._("filetype_docx"), "*.docx"),
                (self._("all_files"), "*.*"),
            ],
            title=self._("export_title")
        )
        if not path:
            return
        ext = os.path.splitext(path)[1].lower()
        try:
            text_lines = self.build_export_text()
            content = '\n'.join(text_lines)
            if ext == ".txt":
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(content)
            elif ext == ".docx":
                if not HAS_DOCX:
                    messagebox.showerror(self._("export_err"),
                        "python-docx is not installed.\nRun: pip install python-docx")
                    return
                doc = Document()
                style = doc.styles['Normal']
                style.font.size = Pt(self.font_size)
                for line in text_lines:
                    doc.add_paragraph(line)
                doc.save(path)
            else:
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(content)
            messagebox.showinfo(self._("export_title"), self._("export_ok").format(path))
        except Exception as e:
            messagebox.showerror(self._("export_err"), str(e))

    # ---- subtitle editor ----
    def _shift_time(self, t, sec):
        try:
            parts = t.replace(',', '.').split(':')
            total = int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2]) + sec
            total = max(0, total)
            h = int(total // 3600)
            m = int((total % 3600) // 60)
            sf = total - h * 3600 - m * 60
            s = int(sf)
            ms = int(round((sf - s) * 1000))
            return f"{h}:{m:02d}:{s:02d}.{ms:03d}"
        except Exception:
            return None

    def _ed_default_times(self):
        ext = os.path.splitext(self.current_filepath or '')[1].lower()
        if ext == '.sub':
            last = self.items[-1].end if self.items else '0'
            try:
                start = str(int(last))
            except Exception:
                start = '0'
            return start, str(int(start) + 25)
        start = self.items[-1].end if self.items else '00:00:00.000'
        end = self._shift_time(start, 2) or '00:00:02.000'
        return start, end

    def open_editor(self):
        if not self.items:
            messagebox.showinfo(self._("edit_title"), self._("edit_no_file"))
            return
        win = getattr(self, "_edit_win", None)
        if win is not None and win.winfo_exists():
            try:
                win.deiconify()
                win.lift()
            except Exception:
                pass
            return
        win = tk.Toplevel(self.root)
        self._edit_win = win
        win.title(self._("edit_title"))
        win.geometry("1350x1368")
        win.minsize(1170, 1152)
        fam = self.font_family
        uifam = self.ui_font_family
        self._ed_style = ttk.Style(win)
        try:
            self._ed_style.configure("UIFont40.TEntry", font=(uifam, 40))
            self._ed_style.configure("UIFont40.TSpinbox", font=(uifam, 40))
        except Exception:
            pass

        tbar = ttk.Frame(win)
        tbar.pack(fill=tk.X, padx=6, pady=6)
        self._edit_tbar = tbar
        for text, cmd in (
            (self._("edit_add"), self._ed_add_item),
            (self._("edit_del"), self._ed_del_item),
            (self._("edit_up"), lambda: self._ed_move(-1)),
            (self._("edit_down"), lambda: self._ed_move(1)),
        ):
            b = ttk.Button(tbar, text=text, command=cmd)
            b.pack(side=tk.LEFT, padx=2)
        ttk.Separator(tbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=6)
        for text, cmd in ((self._("edit_save"), self._ed_save),
                          (self._("edit_save_as"), self._ed_save_as)):
            b = ttk.Button(tbar, text=text, command=cmd)
            b.pack(side=tk.LEFT, padx=2)

        ttk.Separator(tbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=6)
        b = ttk.Button(tbar, text=self._("edit_find_replace"),
                       command=self._open_find_replace)
        b.pack(side=tk.LEFT, padx=2)

        ttk.Separator(tbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=6)
        lbl = ttk.Label(tbar, text=self._("ed_size_list"))
        lbl.pack(side=tk.LEFT, padx=(0, 4))
        self.ed_font_size_var = tk.StringVar(value=str(self.editor_font_size))
        self.ed_font_spin = ttk.Spinbox(tbar, from_=8, to=120, width=8,
                                        textvariable=self.ed_font_size_var,
                                        style="UIFont22.TSpinbox",
                                        command=self._ed_apply_tree_font)
        self.ed_font_spin.configure(font=(self.ui_font_family, 22))
        self.ed_font_spin.pack(side=tk.LEFT, padx=2)
        self.ed_font_size_var.trace_add("write", self._ed_apply_tree_font)
        ttk.Label(tbar, text=self._("ed_size_edit")).pack(side=tk.LEFT, padx=(10, 4))
        self.ed_text_size_var = tk.StringVar(value=str(self.editor_text_size))
        self.ed_text_spin = ttk.Spinbox(tbar, from_=8, to=120, width=8,
                                        textvariable=self.ed_text_size_var,
                                        style="UIFont22.TSpinbox",
                                        command=self._ed_apply_text_font)
        self.ed_text_spin.configure(font=(self.ui_font_family, 22))
        self.ed_text_spin.pack(side=tk.LEFT, padx=2)
        self.ed_text_size_var.trace_add("write", self._ed_apply_text_font)

        paned = ttk.Panedwindow(win, orient=tk.VERTICAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=6, pady=(6, 0))

        top = ttk.Frame(paned)
        colbar = ttk.Frame(top)
        colbar.pack(fill=tk.X, pady=(2, 0))
        ttk.Label(colbar, text=self._("ed_theme")).pack(side=tk.LEFT, padx=(0, 4))
        self._ed_col_btns = {}
        for slot, txt in (("odd_bg", self._("ed_odd_bg")),
                          ("even_bg", self._("ed_even_bg")),
                          ("odd_fg", self._("ed_odd_fg")),
                          ("even_fg", self._("ed_even_fg"))):
            color = getattr(self, "ed_" + slot)
            b = tk.Button(colbar, text=txt, bg=color, fg=self._contrast_color(color),
                          relief=tk.RAISED, bd=1, width=6,
                          font=(uifam, self.ui_font_size),
                          command=lambda s=slot: self._ed_pick_color(s))
            b.pack(side=tk.LEFT, padx=2, pady=2)
            self._ed_col_btns[slot] = b

        wrap = ttk.Frame(top)
        wrap.pack(fill=tk.BOTH, expand=True)
        grid = EdGrid(wrap, self)
        sb = ttk.Scrollbar(wrap, orient=tk.VERTICAL, command=grid.yview)
        hsb = ttk.Scrollbar(wrap, orient=tk.HORIZONTAL, command=grid.hview)
        grid.attach_sb(sb)
        grid.attach_hsb(hsb)
        hsb.pack(side=tk.BOTTOM, fill=tk.X)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        grid.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self._edit_tree = grid
        self._edit_paned = paned
        paned.add(top, weight=4)

        bottom = ttk.Frame(paned)
        form = ttk.Frame(bottom)
        form.pack(fill=tk.X, pady=(4, 2))
        ttk.Label(form, text=self._("edit_start")).grid(row=0, column=0, sticky=tk.W, padx=(0, 4))
        self._ed_start = ttk.Entry(form, width=16, style="UIFont40.TEntry")
        self._ed_start.configure(font=(self.ui_font_family, 40))
        self._ed_start.grid(row=0, column=1, sticky=tk.W, padx=(0, 10))
        ttk.Label(form, text=self._("edit_end")).grid(row=0, column=2, sticky=tk.W, padx=(0, 4))
        self._ed_end = ttk.Entry(form, width=16, style="UIFont40.TEntry")
        self._ed_end.configure(font=(self.ui_font_family, 40))
        self._ed_end.grid(row=0, column=3, sticky=tk.W, padx=(0, 10))
        self._ed_apply_btn = ttk.Button(form, text=self._("edit_apply"), command=self._ed_apply)
        self._ed_apply_btn.grid(row=0, column=4, sticky=tk.W)
        form.columnconfigure(5, weight=1)

        ttk.Label(bottom, text=self._("edit_text")).pack(anchor=tk.W)
        twrap = ttk.Frame(bottom)
        twrap.pack(fill=tk.BOTH, expand=True, pady=(0, 6))
        self._ed_text = tk.Text(twrap, height=6, wrap=tk.WORD,
                                font=(fam, self.editor_text_size))
        tsb = ttk.Scrollbar(twrap, orient=tk.VERTICAL, command=self._ed_text.yview)
        self._ed_text.configure(yscrollcommand=tsb.set)
        self._ed_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        tsb.pack(side=tk.RIGHT, fill=tk.Y)
        paned.add(bottom, weight=2)

        win.protocol("WM_DELETE_WINDOW", self._ed_close)
        self._ed_refresh_tree()
        self._fit_editor_to_work_area(win)

    def _fit_editor_to_work_area(self, win):
        try:
            win.update_idletasks()
            work = get_work_area()
            if not work:
                return
            work_bottom = work["y"] + work["height"]
            work_right = work["x"] + work["width"]
            try:
                w = int(win.winfo_width())
                h = int(win.winfo_height())
            except Exception:
                w = h = 0
            if w <= 1 or h <= 1:
                w = int(win.winfo_reqwidth())
                h = int(win.winfo_reqheight())
            tbar = getattr(self, "_edit_tbar", None)
            if tbar is not None and tbar.winfo_exists():
                try:
                    w = max(w, int(tbar.winfo_reqwidth()) + 30)
                except Exception:
                    pass
            w = min(w, work["width"])
            h = max(200, min(h, work["height"]))
            win.deiconify()
            win.update()
            y = int(win.winfo_y())
            if y + h > work_bottom:
                y = max(work["y"], work_bottom - h)
            x = int(win.winfo_x())
            if x + w > work_right:
                x = max(work["x"], work_right - w)
            win.geometry(f"{w}x{h}+{x}+{y}")
            try:
                mw, mh = win.minsize()
                if h < mh or w < mw:
                    win.minsize(min(w, mw), min(h, mh))
            except Exception:
                pass
            win.lift()
            win.update_idletasks()
            paned = getattr(self, "_edit_paned", None)
            if paned is not None and paned.winfo_exists():
                avail = int(paned.winfo_height())
                if avail > 0:
                    paned.sashpos(0, max(0, int(avail * 4 / 6)))
                paned.update_idletasks()
        except Exception:
            pass

    def _geo_of(self, win):
        try:
            if win is None or not win.winfo_exists():
                return None
            parts = win.geometry().split('+')
            w, h = parts[0].split('x')
            w = int(w)
            h = int(h)
            if w <= 1 or h <= 1:
                w = int(win.winfo_reqwidth())
                h = int(win.winfo_reqheight())
            x = int(parts[1]) if len(parts) > 1 else int(win.winfo_x())
            y = int(parts[2]) if len(parts) > 2 else int(win.winfo_y())
            return x, y, w, h
        except Exception:
            return None

    def _start_display_monitor(self, interval=1500):
        self._display_cache = None
        self._display_interval = interval
        self.root.after(interval, self._poll_display_change)

    def _poll_display_change(self):
        try:
            self._check_display_change()
        except Exception:
            pass
        self.root.after(self._display_interval, self._poll_display_change)

    def _check_display_change(self):
        work = get_work_area()
        scale = get_dpi_scale()
        if not work:
            return
        cache = getattr(self, "_display_cache", None)
        if cache is None:
            self._display_cache = (
                work, scale, self._geo_of(self.root), self._geo_of(getattr(self, "_edit_win", None)))
            return
        old_work, old_scale, old_main, old_edit = cache
        changed = ((old_work["width"], old_work["height"]) != (work["width"], work["height"])
                   or abs(scale - old_scale) > 1e-6)
        if not changed:
            return
        if abs(scale - old_scale) > 1e-6:
            try:
                self.root.tk.call("tk", "scaling", scale)
            except Exception:
                pass
            try:
                self.set_display_style()
            except Exception:
                pass
        fx = work["width"] / max(1, old_work["width"])
        fy = work["height"] / max(1, old_work["height"])
        if old_main:
            try:
                x, y, w, h = old_main
                w = max(200, int(w * fx))
                h = max(200, int(h * fy))
                w = min(w, work["width"])
                h = min(h, work["height"])
                work_bottom = work["y"] + work["height"]
                work_right = work["x"] + work["width"]
                if x + w > work_right:
                    x = max(work["x"], work_right - w)
                if y + h > work_bottom:
                    y = max(work["y"], work_bottom - h)
                self.root.geometry(f"{w}x{h}+{x}+{y}")
            except Exception:
                pass
        win = getattr(self, "_edit_win", None)
        if win is not None and win.winfo_exists() and old_edit:
            try:
                x, y, w, h = old_edit
                w = max(300, int(w * fx))
                h = max(300, int(h * fy))
                w = min(w, work["width"])
                h = min(h, work["height"])
                win.geometry(f"{w}x{h}+{x}+{y}")
                self._fit_editor_to_work_area(win)
            except Exception:
                pass
        self._display_cache = (
            work, scale, self._geo_of(self.root), self._geo_of(win))

    def _ed_close(self):
        win = getattr(self, "_edit_win", None)
        if win is not None and win.winfo_exists():
            win.destroy()
        self._edit_win = None

    def _ed_refresh_tree(self):
        self._edit_tree.invalidate()

    def _contrast_color(self, hexcolor):
        try:
            h = str(hexcolor).lstrip("#")
            r, g, b = (int(h[i:i+2], 16) for i in (0, 2, 4))
            lum = 0.299 * r + 0.587 * g + 0.114 * b
            return "#FFFFFF" if lum < 140 else "#000000"
        except Exception:
            return "#000000"

    def _ed_pick_color(self, slot):
        try:
            key = "ed_" + slot
            current = getattr(self, key)
            from tkinter import colorchooser
            color = colorchooser.askcolor(color=current, parent=self._edit_win,
                                          title=self._("ed_color_title"))
            if not color or not color[1]:
                return
            setattr(self, key, color[1])
            b = self._ed_col_btns.get(slot)
            if b is not None:
                b.configure(bg=color[1], fg=self._contrast_color(color[1]))
            self._ed_refresh_tree()
            self.save_settings()
        except Exception:
            pass

    def _ed_apply_tree_font(self, *args):
        try:
            size = int(self.ed_font_size_var.get())
        except Exception:
            size = self.editor_font_size
        size = max(8, min(120, size))
        self.editor_font_size = size
        self._edit_tree.invalidate()

    def _ed_apply_text_font(self, *args):
        try:
            size = int(self.ed_text_size_var.get())
        except Exception:
            size = self.editor_text_size
        size = max(8, min(120, size))
        self.editor_text_size = size
        fam = self.font_family
        self._ed_text.config(font=(fam, size))

    def _ed_on_select(self, event=None):
        sel = self._edit_tree.selection()
        if not sel:
            return
        try:
            it = self.items[int(sel[0]) - 1]
        except Exception:
            return
        self._ed_start.delete(0, tk.END)
        self._ed_start.insert(0, it.start)
        self._ed_end.delete(0, tk.END)
        self._ed_end.insert(0, it.end)
        self._ed_text.delete(1.0, tk.END)
        self._ed_text.insert(1.0, it.text)

    def _ed_apply(self):
        sel = self._edit_tree.selection()
        if not sel:
            messagebox.showinfo(self._("edit_title"), self._("edit_no_select"))
            return
        start = self._ed_start.get().strip()
        end = self._ed_end.get().strip()
        text = self._ed_text.get(1.0, tk.END).rstrip('\n')
        ext = os.path.splitext(self.current_filepath or '')[1].lower()
        if ext == '.sub':
            ok = bool(re.match(r'^\d+$', start)) and bool(re.match(r'^\d+$', end))
        else:
            ok = bool(SubtitleReader._TIME_RE.match(start)
                      and SubtitleReader._TIME_RE.match(end))
        if not ok:
            messagebox.showerror(self._("edit_title"),
                                 self._("edit_invalid_time").format(f"{start} -> {end}"))
            return
        for tid in sel:
            try:
                it = self.items[int(tid) - 1]
                it.start = start.replace(',', '.')
                it.end = end.replace(',', '.')
                it.text = text
                it.duration = SubtitleParser._calc_duration(it.start, it.end)
            except Exception:
                pass
        self._ed_refresh_tree()
        self._edit_tree.selection_set(sel)

    def _open_find_replace(self):
        win = getattr(self, "_edit_win", None)
        if win is None or not win.winfo_exists():
            return
        if getattr(self, "_fr_win", None) is not None and self._fr_win.winfo_exists():
            try:
                self._fr_win.deiconify()
                self._fr_win.lift()
                self._fr_find.focus_set()
            except Exception:
                pass
            return
        dlg = tk.Toplevel(win)
        self._fr_win = dlg
        dlg.title(self._("edit_find_title"))
        dlg.transient(win)
        uifam = self.ui_font_family
        fsz = self.ui_font_size
        frame = ttk.Frame(dlg)
        frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        frame.columnconfigure(1, weight=1)
        ttk.Label(frame, text=self._("edit_find_what")).grid(row=0, column=0, sticky="e",
                                                             padx=4, pady=3)
        self._fr_find = tk.Entry(frame, font=(uifam, fsz))
        self._fr_find.grid(row=0, column=1, columnspan=3, sticky="ew", padx=4, pady=3)
        ttk.Label(frame, text=self._("edit_replace_with")).grid(row=1, column=0, sticky="e",
                                                                padx=4, pady=3)
        self._fr_rep = tk.Entry(frame, font=(uifam, fsz))
        self._fr_rep.grid(row=1, column=1, columnspan=3, sticky="ew", padx=4, pady=3)
        self._fr_time = tk.BooleanVar(value=False)
        self._fr_case = tk.BooleanVar(value=False)
        ttk.Checkbutton(frame, text=self._("edit_find_time"),
                        variable=self._fr_time).grid(row=2, column=0, columnspan=2,
                                                     sticky="w", padx=4)
        ttk.Checkbutton(frame, text=self._("edit_find_case"),
                        variable=self._fr_case).grid(row=2, column=2, columnspan=2,
                                                     sticky="w", padx=4)
        self._fr_status = ttk.Label(frame, text="")
        self._fr_status.grid(row=3, column=0, columnspan=4, sticky="w", padx=4, pady=(2, 4))
        row = 4
        for col, (text, cmd) in enumerate((("edit_find_prev", self._fr_prev),
                                           ("edit_find_next", self._fr_next),
                                           ("edit_find_replace_one", self._fr_replace_one),
                                           ("edit_find_replace_all", self._fr_replace_all))):
            b = ttk.Button(frame, text=self._(text), command=cmd)
            b.grid(row=row, column=col, sticky="ew", padx=4, pady=3)
        for c in range(4):
            frame.columnconfigure(c, weight=1)
        self._fr_pos = -1
        self._fr_matches = []
        dlg.protocol("WM_DELETE_WINDOW", self._fr_close)
        try:
            sel = self._edit_tree.selection()
            if sel:
                self._fr_find.insert(0, self.items[int(sel[0]) - 1].text.split("\n")[0])
        except Exception:
            pass
        self._fr_find.focus_set()
        self._fr_find.select_range(0, tk.END)

    def _fr_close(self):
        try:
            self._fr_win.destroy()
        except Exception:
            pass
        self._fr_win = None

    def _fr_contains(self, v, q, case):
        if not v:
            return False
        return (q in v) if case else (q.lower() in v.lower())

    def _fr_matches_of(self, v, q, case):
        if not v or not q:
            return 0
        return v.count(q) if case else v.lower().count(q.lower())

    def _fr_replace_in(self, v, q, rep, case):
        if not v or not q:
            return v
        if case:
            return v.replace(q, rep)
        try:
            import re as _re
            return _re.sub(_re.escape(q), lambda mm: rep, v, flags=_re.IGNORECASE)
        except Exception:
            low = v.lower()
            ql = q.lower()
            out = []
            i = 0
            while True:
                j = low.find(ql, i)
                if j < 0:
                    out.append(v[i:])
                    break
                out.append(v[i:j])
                out.append(rep)
                i = j + len(q)
            return "".join(out)

    def _fr_gather(self):
        self._fr_matches = []
        self._fr_pos = -1
        q = self._fr_find.get()
        if not q:
            self._fr_status.config(text=self._("edit_find_no_text"))
            return False
        find_time = self._fr_time.get()
        case = self._fr_case.get()
        for i, it in enumerate(self.items):
            if find_time:
                for k in ("text", "start", "end"):
                    if self._fr_contains(getattr(it, k, ""), q, case):
                        self._fr_matches.append((i, k))
            else:
                if self._fr_contains(it.text, q, case):
                    self._fr_matches.append((i, "text"))
        self._fr_status.config(text=self._("edit_find_found").format(len(self._fr_matches)))
        if not self._fr_matches:
            self._fr_status.config(text=self._("edit_find_not_found").format(q))
            self._edit_tree.selection_set([])
            return False
        return True

    def _fr_reveal(self, i):
        grid = self._edit_tree
        try:
            rh = grid._rowh()
            hh = grid._headh()
            H = grid.winfo_height()
            grid._off = max(0, i * rh - int((H - hh) * 0.35))
            grid.redraw()
        except Exception:
            pass
        self._edit_tree.selection_set(str(i + 1))
        self._ed_on_select()

    def _fr_next(self):
        if not self._fr_gather():
            self._fr_find.focus_set()
            return
        self._fr_pos = (self._fr_pos + 1) % len(self._fr_matches)
        i, _ = self._fr_matches[self._fr_pos]
        self._fr_reveal(i)

    def _fr_prev(self):
        if not self._fr_gather():
            self._fr_find.focus_set()
            return
        self._fr_pos = (self._fr_pos - 1) % len(self._fr_matches)
        i, _ = self._fr_matches[self._fr_pos]
        self._fr_reveal(i)

    def _fr_replace_one(self):
        if not self._fr_gather():
            self._fr_find.focus_set()
            return
        q = self._fr_find.get()
        rep = self._fr_rep.get()
        case = self._fr_case.get()
        self._fr_pos = (self._fr_pos + 1) % len(self._fr_matches)
        i, k = self._fr_matches[self._fr_pos]
        it = self.items[i]
        cur = getattr(it, k, "") or ""
        setattr(it, k, self._fr_replace_in(cur, q, rep, case))
        self._fr_reveal(i)
        self._fr_gather()

    def _fr_replace_all(self):
        if not self._fr_gather():
            self._fr_find.focus_set()
            return
        q = self._fr_find.get()
        rep = self._fr_rep.get()
        case = self._fr_case.get()
        find_time = self._fr_time.get()
        total = 0
        for i, it in enumerate(self.items):
            if find_time:
                for k in ("text", "start", "end"):
                    cur = getattr(it, k, "") or ""
                    n = self._fr_matches_of(cur, q, case)
                    if n:
                        setattr(it, k, self._fr_replace_in(cur, q, rep, case))
                        total += n
            else:
                cur = it.text or ""
                n = self._fr_matches_of(cur, q, case)
                if n:
                    it.text = self._fr_replace_in(cur, q, rep, case)
                    total += n
        self._fr_status.config(text=self._("edit_find_replaced").format(total))
        self._fr_matches = []
        self._fr_pos = -1
        self._ed_refresh_tree()

    def _ed_add_item(self):
        start, end = self._ed_default_times()
        self.items.append(SubtitleItem(len(self.items) + 1, start, end,
                                       self._("edit_new_text"), ""))
        self._ed_refresh_tree()
        self._edit_tree.selection_set(str(len(self.items)))
        self._ed_on_select()

    def _ed_del_item(self):
        sel = self._edit_tree.selection()
        if not sel:
            messagebox.showinfo(self._("edit_title"), self._("edit_no_select"))
            return
        if not messagebox.askyesno(self._("edit_confirm"),
                                   self._("edit_del_confirm").format(len(sel))):
            return
        for tid in sorted((int(t) - 1 for t in sel), reverse=True):
            if 0 <= tid < len(self.items):
                self.items.pop(tid)
        self._ed_refresh_tree()

    def _ed_move(self, delta):
        sel = self._edit_tree.selection()
        if not sel:
            return
        idcs = sorted(int(t) - 1 for t in sel)
        lo, hi = idcs[0], idcs[-1]
        if delta < 0:
            if lo <= 0:
                return
            block = self.items[lo:hi + 1]
            del self.items[lo:hi + 1]
            self.items[lo - 1:lo - 1] = block
        else:
            if hi >= len(self.items) - 1:
                return
            block = self.items[lo:hi + 1]
            del self.items[lo:hi + 1]
            self.items[lo + 1:lo + 1] = block
        self._ed_refresh_tree()
        new_ids = [str(i - 1) for i in idcs] if delta < 0 else [str(i + 1) for i in idcs]
        self._edit_tree.selection_set(new_ids)
        self._ed_on_select()

    def _renumber_items(self):
        for i, it in enumerate(self.items, 1):
            it.index = i

    def _refresh_after_edit(self):
        self._renumber_items()
        try:
            self._ed_refresh_tree()
            self._edit_tree.selection_set(str(len(self.items)))
        except Exception:
            pass
        self.refresh_display()
        self.refresh_status_bar()
        self.update_progress()

    def _write_items(self, path):
        try:
            SubtitleParser.write_file(path, self.items)
            messagebox.showinfo(self._("edit_title"), self._("edit_saved").format(path))
            return True
        except Exception as e:
            messagebox.showerror(self._("edit_title"), self._("edit_save_err").format(e))
            return False

    def _ed_save(self):
        if not self.current_filepath:
            self._ed_save_as()
            return
        if not messagebox.askyesno(self._("edit_confirm"),
                                   self._("edit_overwrite").format(self.current_filepath)):
            return
        if self._write_items(self.current_filepath):
            self._refresh_after_edit()

    def _ed_save_as(self):
        cur = self.current_filepath or ""
        base = os.path.splitext(os.path.basename(cur))[0] if cur else "subtitles"
        ext = os.path.splitext(cur)[1].lower() or ".srt"
        if ext not in SUBTITLE_EXTS:
            ext = ".srt"
        path = filedialog.asksaveasfilename(
            initialfile=base,
            defaultextension=ext,
            filetypes=[(self._("file_filter"), "*.srt *.sub *.ass *.ssa"),
                       (self._("all_files"), "*.*")],
            title=self._("edit_save_as")
        )
        if not path:
            return
        if os.path.splitext(path)[1].lower() not in SUBTITLE_EXTS:
            path += ext
        if self._write_items(path):
            self.current_filepath = path
            self.root.title(f"{self._('app_title')} - {os.path.basename(path)}")
            self.update_recent_files(path)
            self.rebuild_file_menu()
            self._refresh_after_edit()

    # ---- config ----
    def get_config_dir(self):
        if sys.platform == "win32":
            base = os.environ.get("APPDATA", os.path.expanduser("~"))
        else:
            base = os.path.expanduser("~/.config")
        cfg_dir = os.path.join(base, "SubtitleReader")
        os.makedirs(cfg_dir, exist_ok=True)
        return cfg_dir

    def get_config_path(self):
        return os.path.join(self.get_config_dir(), "config.json")

    def load_config(self):
        try:
            with open(self.get_config_path(), encoding="utf-8") as f:
                return json.load(f)
        except:
            return {"recent_files": [], "positions": {}}

    def save_config(self, cfg):
        try:
            with open(self.get_config_path(), "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False, indent=2)
        except:
            pass

    def update_recent_files(self, filepath):
        cfg = self.load_config()
        recent = cfg.get("recent_files", [])
        if filepath in recent:
            recent.remove(filepath)
        recent.insert(0, filepath)
        cfg["recent_files"] = recent[:10]
        self.save_config(cfg)

    # ---- position ----
    def get_current_index(self):
        if not self.items:
            return 0
        try:
            frac = self.text_area.yview()[0]
            return max(0, min(int(frac * len(self.items)), len(self.items) - 1))
        except:
            return 0

    def save_position(self):
        if not self.current_filepath or not self.items:
            return
        cfg = self.load_config()
        if "positions" not in cfg:
            cfg["positions"] = {}
        idx = self.get_current_index()
        frac = self.text_area.yview()[0]
        cfg["positions"][self.current_filepath] = {"index": idx, "fraction": frac}
        self.save_config(cfg)

    def restore_position(self):
        if not self.current_filepath:
            return
        cfg = self.load_config()
        pos = cfg.get("positions", {}).get(self.current_filepath)
        if pos:
            frac = pos.get("fraction", 0.0)
            self.root.after_idle(lambda: self.text_area.yview_moveto(frac))

    # ---- close ----
    def close_file(self):
        if not self.items:
            return
        self.save_position()
        if getattr(self, "_edit_win", None) is not None and self._edit_win.winfo_exists():
            self._ed_close()
        self.items = []
        self.current_filepath = None
        self.root.title(self._("app_title"))
        self.refresh_display()
        self.refresh_status_bar()
        self.update_progress()
        self.rebuild_file_menu()

    # ---- recent files ----
    def open_recent_file(self, filepath):
        if not os.path.exists(filepath):
            cfg = self.load_config()
            recent = cfg.get("recent_files", [])
            if filepath in recent:
                recent.remove(filepath)
            cfg["recent_files"] = recent
            self.save_config(cfg)
            self.rebuild_file_menu()
            return
        self.open_file(filepath=filepath, restore_pos=True)

    # ---- progress ----
    def update_progress(self, *args):
        if not self.items:
            self.progress_bar.configure(maximum=100)
            self.progress_bar["value"] = 0
            self.progress_label.config(text="0%")
            return
        try:
            if self.display_mode.get() == "percent":
                frac = self.text_area.yview()[0]
                val = int(frac * 100)
                self.progress_bar.configure(maximum=100)
                self.progress_bar["value"] = val
                self.progress_label.config(text=f"{val}%")
            else:
                yv = self.text_area.yview()
                visible = yv[1] - yv[0]
                if visible <= 0:
                    visible = 0.01
                total = max(1, int(1.0 / visible + 0.5))
                cur = int(yv[0] * total) + 1
                cur = max(1, min(cur, total))
                self.progress_bar.configure(maximum=total)
                self.progress_bar["value"] = cur
                self.progress_label.config(text=self._("page_fmt").format(cur, total))
        except:
            self.progress_bar.configure(maximum=100)
            self.progress_bar["value"] = 0
            self.progress_label.config(text="0%")

    def _on_scrollbar(self, *args):
        self.text_area.yview(*args)
        self.update_progress()

    def _on_text_scroll(self, *args):
        self.scrollbar.set(*args)
        self.update_progress()

    def _on_mousewheel(self, event):
        self.root.after(10, self.update_progress)

    # ---- menu rebuild ----
    def rebuild_file_menu(self):
        self.file_menu.delete(0, "end")
        has_file = bool(self.items)

        self.w["edit_btn"].config(state=tk.NORMAL if has_file else tk.DISABLED)

        self.file_menu.add_command(label=self._("menu_open"), command=self.open_file, accelerator="Ctrl+O")
        self.file_menu.add_command(
            label=self._("menu_edit"), command=self.open_editor,
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_command(
            label=self._("menu_close"), command=self.close_file,
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_separator()
        self.file_menu.add_command(
            label=self._("menu_export"), command=self.export_content, accelerator="Ctrl+E",
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_separator()

        cfg = self.load_config()
        recent = cfg.get("recent_files", [])
        positions = cfg.get("positions", {})
        self._last_rebuild_recent = list(recent)
        if recent:
            self.file_menu.add_command(
                label="── " + self._("menu_recent") + " ──",
                state=tk.DISABLED
            )
            for fpath in recent:
                exists = os.path.exists(fpath)
                label = os.path.basename(fpath)
                pos = positions.get(fpath, {})
                frac = pos.get("fraction", 0.0)
                pct = min(100, max(0, int(frac * 100)))
                if 0 < pct < 100:
                    label = f"[{pct}%] {label}"
                elif pct >= 100:
                    label = f"[100%] {label}"
                if not exists:
                    label += " " + self._("file_missing")
                if exists:
                    self.file_menu.add_command(
                        label=label,
                        command=lambda p=fpath: self.open_recent_file(p)
                    )
                else:
                    self.file_menu.add_command(label=label, state=tk.DISABLED)
            self.file_menu.add_separator()
            self.file_menu.add_command(
                label=self._("clear_recent"),
                command=self.clear_recent_files
            )

        self.file_menu.add_separator()
        self.file_menu.add_command(label=self._("menu_exit"), command=self.root.quit)

    def clear_recent_files(self):
        if not messagebox.askyesno(self._("clear_recent"), self._("clear_recent_confirm")):
            return
        cfg = self.load_config()
        cfg["recent_files"] = []
        self.save_config(cfg)
        self.rebuild_file_menu()

    def _on_progress_click(self, event):
        if not self.items:
            return
        width = self.progress_bar.winfo_width()
        if width <= 0:
            return
        fraction = event.x / width
        if self.display_mode.get() == "percent":
            target_idx = int(fraction * len(self.items))
            target_idx = max(0, min(target_idx, len(self.items) - 1))
            target_frac = target_idx / len(self.items)
            self.text_area.yview_moveto(target_frac)
        else:
            yv = self.text_area.yview()
            visible = yv[1] - yv[0]
            if visible <= 0:
                visible = 0.01
            total = max(1, int(1.0 / visible + 0.5))
            target_page = int(fraction * total) + 1
            target_page = max(1, min(target_page, total))
            target_frac = (target_page - 1) / total
            self.text_area.yview_moveto(target_frac)
        self.update_progress()

    def _on_progress_label_click(self, event):
        if not self.items:
            return
        if self.display_mode.get() == "percent":
            cur_val = self.progress_bar["value"]
            result = simpledialog.askinteger(
                self._("jump_to"), self._("jump_to_hint"),
                minvalue=0, maxvalue=100,
                initialvalue=int(cur_val),
                parent=self.root
            )
            if result is not None:
                fraction = result / 100.0
                target_idx = int(fraction * len(self.items))
                target_idx = max(0, min(target_idx, len(self.items) - 1))
                target_frac = target_idx / len(self.items)
                self.text_area.yview_moveto(target_frac)
        else:
            yv = self.text_area.yview()
            visible = yv[1] - yv[0]
            if visible <= 0:
                visible = 0.01
            total = max(1, int(1.0 / visible + 0.5))
            cur_page = int(yv[0] * total) + 1
            cur_page = max(1, min(cur_page, total))
            result = simpledialog.askinteger(
                self._("jump_to"), self._("page_jump_hint").format(total),
                minvalue=1, maxvalue=total,
                initialvalue=cur_page,
                parent=self.root
            )
            if result is not None:
                target_frac = (result - 1) / total
                self.text_area.yview_moveto(target_frac)
        self.update_progress()

    def on_display_mode_change(self):
        self.update_progress()

    def show_about(self):
        win = tk.Toplevel(self.root)
        win.title(self._("about_title"))
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        lines = self._("about_info").split('\n')
        for i, line in enumerate(lines):
            lbl = tk.Label(win, text=line, font=(self.font_family, self.ui_font_size), anchor=tk.CENTER)
            lbl.pack(fill=tk.X, padx=30, pady=(10 if i == 0 else 0, 0))
        tk.Button(win, text=self._("menu_close"), command=win.destroy,
                  width=10, font=(self.font_family, self.ui_font_size)).pack(pady=15)
        win.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() - win.winfo_width()) // 2
        y = self.root.winfo_rooty() + 50
        win.geometry(f"+{x}+{y}")

    def set_window_icon(self):
        base = getattr(sys, '_MEIPASS', os.path.dirname(__file__))
        ico_path = os.path.join(base, "icon.ico")
        png_path = os.path.join(base, "icon.png")
        if os.path.exists(ico_path):
            try:
                self.root.iconbitmap(ico_path)
            except:
                pass
        elif os.path.exists(png_path):
            try:
                img = tk.PhotoImage(file=png_path)
                self.root.iconphoto(True, img)
            except:
                pass

if __name__ == "__main__":
    enable_high_dpi()
    dnd_root = None
    if HAS_DND:
        try:
            dnd_root = TkinterDnD.Tk()
        except Exception:
            HAS_DND = False
            dnd_root = None
    root = dnd_root if dnd_root is not None else tk.Tk()
    scale = get_dpi_scale()
    try:
        root.tk.call("tk", "scaling", scale if scale != 1.0 else 1.0)
    except Exception:
        pass
    work = get_work_area()
    width = int(1200 * scale) if scale != 1.0 else 1200
    if work and work["height"] > 0:
        max_w = work["width"]
        if width > max_w:
            width = max_w
        x = work["x"] + (max_w - width) // 2
        root.withdraw()
        root.geometry(f"{width}x{work['height']}+{x}+{work['y']}")
        root.update_idletasks()
        root.deiconify()
        root.update()
    else:
        if scale != 1.0:
            root.geometry(f"{int(1200 * scale)}x{int(1200 * scale)}")
        else:
            root.geometry("1200x1200")
    app = SubtitleReader(root)
    if work and work["height"] > 0:
        root.update_idletasks()
        try:
            toolbar_w = app.toolbar.winfo_reqwidth() + 30
            if toolbar_w > width:
                ov = toolbar_w - width
                if width + ov <= work["width"]:
                    width += ov
                else:
                    width = work["width"]
                x = work["x"] + (max_w - width) // 2
                root.geometry(f"{width}x{work['height']}+{x}+{work['y']}")
                root.update_idletasks()
            rect = _window_outer(root)
            if rect:
                overflow = (rect[1] + rect[3]) - (work["y"] + work["height"])
                if overflow > 0:
                    height = max(200, work["height"] - overflow)
                    root.geometry(f"{width}x{height}+{x}+{work['y']}")
                    root.update_idletasks()
        except Exception:
            pass
    app._start_display_monitor()
    root.mainloop()
