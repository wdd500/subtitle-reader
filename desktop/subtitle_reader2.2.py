import sys
import ctypes
import tkinter as tk
from tkinter import filedialog, colorchooser, font, ttk, messagebox, simpledialog
import re
import os
import json

try:
    from docx import Document
    from docx.shared import Pt
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except ImportError:
    HAS_DND = False


LANG = {
    "zh": {
        "app_title": "字幕阅读器",
        "menu_file": "文件",
        "menu_open": "打开...",
        "menu_export": "导出...",
        "menu_exit": "退出",
        "menu_lang": "语言",
        "lang_zh": "中文",
        "lang_en": "English",
        "open_file": "打开文件",
        "timecode": "时间码",
        "line_no": "行号",
        "duration": "时长",
        "font": "字体：",
        "size": "字号：",
        "text_color": "文字颜色",
        "bg_color": "背景颜色",
        "margin": "边距：",
        "status_no_file": "未加载文件",
        "status_items": "{} 条字幕",
        "status_duration": "总时长：{}",
        "no_file": "未加载字幕。\n打开或将 .srt/.sub/.ass 文件拖入窗口开始阅读。",
        "export_title": "导出字幕内容",
        "export_ok": "已导出到：\n{}",
        "export_err": "导出错误",
        "export_none": "未加载字幕文件。",
        "text_color_title": "文字颜色",
        "bg_color_title": "背景颜色",
        "sub_count": "{} 条字幕",
        "open_err": "打开文件出错：{}",
        "file_filter": "字幕文件",
        "all_files": "所有文件",
        "filetype_txt": "文本文件",
        "filetype_docx": "Word 文档",
        "export_fmt": "导出格式",
        "export_txt": "文本 (.txt)",
        "export_docx": "Word (.docx)",
        "menu_close": "关闭",
        "menu_recent": "最近打开",
        "file_missing": "(文件不存在)",
        "clear_recent": "清除记录",
        "clear_recent_confirm": "确定清除所有最近打开记录？",
        "jump_to": "跳转到位置",
        "jump_to_hint": "请输入百分比 (0-100)：",
        "menu_display": "显示模式",
        "mode_percent": "百分比%",
        "mode_page": "分页",
        "page_fmt": "第{}页/共{}页",
        "page_jump_hint": "请输入页数 (1-{})：",
        "menu_help": "帮助",
        "menu_about": "关于",
        "about_title": "关于字幕阅读器",
        "about_info": "字幕阅读器 v2.2\nSubtitle-Reader v2.2\n\n作者：武东东\n开源地址\nhttps://github.com/wdd500/subtitle-reader",
    },
    "en": {
        "app_title": "Subtitle Reader",
        "menu_file": "File",
        "menu_open": "Open...",
        "menu_export": "Export...",
        "menu_exit": "Exit",
        "menu_lang": "Language",
        "lang_zh": "中文",
        "lang_en": "English",
        "open_file": "Open File",
        "timecode": "Timecode",
        "line_no": "Line No.",
        "duration": "Duration",
        "font": "Font:",
        "size": "Size:",
        "text_color": "Text Color",
        "bg_color": "Bg Color",
        "margin": "Margin:",
        "status_no_file": "No file loaded",
        "status_items": "{} subtitles",
        "status_duration": "Total: {}",
        "no_file": "No subtitles loaded.\nOpen or drop a .srt / .sub / .ass file to begin.",
        "export_title": "Export Subtitle Content",
        "export_ok": "Exported to:\n{}",
        "export_err": "Export Error",
        "export_none": "No subtitle file loaded.",
        "text_color_title": "Text Color",
        "bg_color_title": "Background Color",
        "sub_count": "{} subtitles",
        "open_err": "Error opening file: {}",
        "file_filter": "Subtitle files",
        "all_files": "All files",
        "filetype_txt": "Text file",
        "filetype_docx": "Word document",
        "export_fmt": "Export Format",
        "export_txt": "Text (.txt)",
        "export_docx": "Word (.docx)",
        "menu_close": "Close",
        "menu_recent": "Recent Files",
        "file_missing": "(file missing)",
        "clear_recent": "Clear Recent Files",
        "clear_recent_confirm": "Clear all recent files?",
        "jump_to": "Jump to Position",
        "jump_to_hint": "Enter percentage (0-100):",
        "menu_display": "Display Mode",
        "mode_percent": "Percentage %",
        "mode_page": "Paging",
        "page_fmt": "Page {}/{}",
        "page_jump_hint": "Enter page (1-{}):",
        "menu_help": "Help",
        "menu_about": "About",
        "about_title": "About Subtitle Reader",
        "about_info": "字幕阅读器 v2.2\nSubtitle-Reader v2.2\n\n作者：武东东\n开源地址\nhttps://github.com/wdd500/subtitle-reader",
    },
}


def enable_high_dpi():
    if sys.platform != "win32":
        return
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


def get_dpi_scale():
    try:
        hwnd = ctypes.windll.user32.GetDesktopWindow()
        dpi = ctypes.windll.user32.GetDpiForWindow(hwnd)
        return dpi / 96.0
    except Exception:
        return 1.0


def get_work_area():
    class RECT(ctypes.Structure):
        _fields_ = [("left", ctypes.c_long), ("top", ctypes.c_long),
                    ("right", ctypes.c_long), ("bottom", ctypes.c_long)]
    rect = RECT()
    try:
        ok = ctypes.windll.user32.SystemParametersInfoW(
            0x0030, 0, ctypes.byref(rect), 0)
        if ok:
            return {
                "x": rect.left,
                "y": rect.top,
                "width": rect.right - rect.left,
                "height": rect.bottom - rect.top,
            }
    except Exception:
        pass
    return None


def _window_outer(root):
    class RECT(ctypes.Structure):
        _fields_ = [("left", ctypes.c_long), ("top", ctypes.c_long),
                    ("right", ctypes.c_long), ("bottom", ctypes.c_long)]
    try:
        hwnd = root.winfo_id()
        rect = RECT()
        if ctypes.windll.user32.GetWindowRect(hwnd, ctypes.byref(rect)):
            return rect.left, rect.top, rect.right - rect.left, rect.bottom - rect.top
    except Exception:
        pass
    return None


SUBTITLE_EXTS = ('.srt', '.sub', '.ass', '.ssa')


class SubtitleItem:
    def __init__(self, index=0, start="", end="", text="", duration=""):
        self.index = index
        self.start = start
        self.end = end
        self.text = text
        self.duration = duration


class SubtitleParser:
    @staticmethod
    def parse_srt(content):
        items = []
        blocks = re.split(r'\n\s*\n', content.strip())
        for block in blocks:
            lines = block.strip().split('\n')
            if len(lines) >= 3:
                try:
                    idx = int(lines[0])
                    time_match = re.match(
                        r'(\d+:\d+:\d+[,.]\d+)\s*-->\s*(\d+:\d+:\d+[,.]\d+)', lines[1])
                    if time_match:
                        start = time_match.group(1).replace(',', '.')
                        end = time_match.group(2).replace(',', '.')
                        text = '\n'.join(lines[2:])
                        duration = SubtitleParser._calc_duration(start, end)
                        items.append(SubtitleItem(idx, start, end, text, duration))
                except ValueError:
                    continue
        return items

    @staticmethod
    def parse_sub(content):
        items = []
        lines = content.strip().split('\n')
        for i, line in enumerate(lines):
            line = line.strip()
            if not line:
                continue
            match = re.match(r'\{(\d+)\}\{(\d+)\}(.*)', line)
            if match:
                start_frame = match.group(1)
                end_frame = match.group(2)
                text = match.group(3).replace('|', '\n')
                duration = str(int(end_frame) - int(start_frame)) + ' frames'
                items.append(SubtitleItem(i + 1, start_frame, end_frame, text, duration))
        return items

    @staticmethod
    def parse_ass(content):
        items = []
        events_section = False
        format_line = None
        for line in content.split('\n'):
            line = line.strip()
            if line.startswith('[Events]'):
                events_section = True
                continue
            if events_section:
                if line.startswith('Format:'):
                    format_line = line[7:].strip().split(',')
                    format_line = [f.strip() for f in format_line]
                    continue
                if line.startswith('Dialogue:'):
                    parts = line[9:].strip()
                    if format_line:
                        vals = SubtitleParser._split_ass(parts)
                        data = {}
                        for j, key in enumerate(format_line):
                            if j < len(vals):
                                data[key] = vals[j]
                        start = data.get('Start', '')
                        end = data.get('End', '')
                        text = data.get('Text', '')
                        text = re.sub(r'\{[^}]*\}', '', text)
                        text = text.replace('\\N', '\n').replace('\\n', '\n')
                        duration = SubtitleParser._calc_duration_ass(start, end)
                        items.append(SubtitleItem(len(items) + 1, start, end, text, duration))
        return items

    @staticmethod
    def _split_ass(line):
        parts = []
        current = ''
        in_braces = 0
        for ch in line:
            if ch == '{':
                in_braces += 1
                current += ch
            elif ch == '}':
                in_braces -= 1
                current += ch
            elif ch == ',' and in_braces == 0:
                parts.append(current.strip())
                current = ''
            else:
                current += ch
        parts.append(current.strip())
        return parts

    @staticmethod
    def _calc_duration(start, end):
        def to_seconds(t):
            parts = t.replace(',', '.').split(':')
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
        try:
            d = to_seconds(end) - to_seconds(start)
            return f'{d:.2f}s'
        except:
            return ''

    @staticmethod
    def _calc_duration_ass(start, end):
        def to_seconds(t):
            parts = t.replace(',', '.').split(':')
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
        try:
            d = to_seconds(end) - to_seconds(start)
            return f'{d:.2f}s'
        except:
            return ''

    @staticmethod
    def parse_file(filepath):
        ext = os.path.splitext(filepath)[1].lower()
        with open(filepath, encoding='utf-8') as f:
            content = f.read()
        if ext == '.srt':
            return SubtitleParser.parse_srt(content)
        elif ext == '.sub':
            return SubtitleParser.parse_sub(content)
        elif ext in ('.ass', '.ssa'):
            return SubtitleParser.parse_ass(content)
        else:
            return []


class SubtitleReader:
    def __init__(self, root):
        self.root = root
        self.current_lang = "zh"
        self.items = []
        self.current_filepath = None
        self.show_time = False
        self.show_index = False
        self.show_duration = False
        self.text_color = "#000000"
        self.bg_color = "#FFFFFF"
        self.font_size = 22
        self.font_family = "Segoe UI"
        self.ui_font_size = 18
        self.margin = 10
        self.w = {}
        self._last_rebuild_recent = []
        self.display_mode = tk.StringVar(value="percent")
        self.build_ui()
        self.apply_language()
        self.load_settings()
        if HAS_DND:
            self.register_drop_targets(self.root)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---- i18n helpers ----
    def _(self, key):
        return LANG[self.current_lang].get(key, key)

    def apply_language(self):
        self.root.title(self._("app_title"))
        self.rebuild_file_menu()

        self.lang_menu.delete(0, "end")
        self.lang_menu.add_command(label=self._("lang_zh"), command=lambda: self.set_language("zh"))
        self.lang_menu.add_command(label=self._("lang_en"), command=lambda: self.set_language("en"))

        self.display_menu.delete(0, "end")
        self.display_menu.add_radiobutton(label=self._("mode_percent"), variable=self.display_mode, value="percent", command=self.on_display_mode_change)
        self.display_menu.add_radiobutton(label=self._("mode_page"), variable=self.display_mode, value="page", command=self.on_display_mode_change)

        self.help_menu.delete(0, "end")
        self.help_menu.add_command(label=self._("menu_about"), command=self.show_about)

        self.menubar.delete(0, "end")
        self.menubar.add_cascade(label=self._("menu_file"), menu=self.file_menu)
        self.menubar.add_cascade(label=self._("menu_display"), menu=self.display_menu)
        self.menubar.add_cascade(label=self._("menu_lang"), menu=self.lang_menu)
        self.menubar.add_cascade(label=self._("menu_help"), menu=self.help_menu)
        self.root.config(menu=self.menubar)

        self.w["open_btn"].config(text=self._("open_file"))
        self.w["time_cb"].config(text=self._("timecode"))
        self.w["idx_cb"].config(text=self._("line_no"))
        self.w["dur_cb"].config(text=self._("duration"))
        self.w["font_lbl"].config(text=self._("font"))
        self.w["size_lbl"].config(text=self._("size"))
        self.w["txt_color_btn"].config(text=self._("text_color"))
        self.w["bg_color_btn"].config(text=self._("bg_color"))
        self.w["margin_lbl"].config(text=self._("margin"))

        self.refresh_display()
        self.refresh_status_bar()

    # ---- build UI ----
    def build_ui(self):
        self.menubar = tk.Menu(self.root)
        self.file_menu = tk.Menu(self.menubar, tearoff=0)
        self.file_menu.add_command(label="", command=self.open_file, accelerator="Ctrl+O")
        self.file_menu.add_separator()
        self.file_menu.add_command(label="", command=self.export_content, accelerator="Ctrl+E")
        self.file_menu.add_separator()
        self.file_menu.add_command(label="", command=self.root.quit)
        self.menubar.add_cascade(label="", menu=self.file_menu)

        self.display_menu = tk.Menu(self.menubar, tearoff=0)
        self.display_menu.add_radiobutton(label="", variable=self.display_mode, value="percent", command=self.on_display_mode_change)
        self.display_menu.add_radiobutton(label="", variable=self.display_mode, value="page", command=self.on_display_mode_change)
        self.menubar.add_cascade(label="", menu=self.display_menu)

        self.lang_menu = tk.Menu(self.menubar, tearoff=0)
        self.lang_menu.add_command(label="", command=lambda: self.set_language("zh"))
        self.lang_menu.add_command(label="", command=lambda: self.set_language("en"))
        self.menubar.add_cascade(label="", menu=self.lang_menu)

        self.help_menu = tk.Menu(self.menubar, tearoff=0)
        self.help_menu.add_command(label="", command=self.show_about)
        self.menubar.add_cascade(label="", menu=self.help_menu)

        self.root.config(menu=self.menubar)
        self.set_window_icon()
        self.root.bind_all("<Control-o>", lambda e: self.open_file())
        self.root.bind_all("<Control-e>", lambda e: self.export_content())

        toolbar = ttk.Frame(self.root)
        toolbar.pack(fill=tk.X, padx=3, pady=3)

        self.w["open_btn"] = ttk.Button(toolbar, command=self.open_file)
        self.w["open_btn"].pack(side=tk.LEFT, padx=2)

        self.time_var = tk.BooleanVar()
        self.w["time_cb"] = ttk.Checkbutton(toolbar, variable=self.time_var, command=self.refresh_display)
        self.w["time_cb"].pack(side=tk.LEFT, padx=2)

        self.index_var = tk.BooleanVar()
        self.w["idx_cb"] = ttk.Checkbutton(toolbar, variable=self.index_var, command=self.refresh_display)
        self.w["idx_cb"].pack(side=tk.LEFT, padx=2)

        self.dur_var = tk.BooleanVar()
        self.w["dur_cb"] = ttk.Checkbutton(toolbar, variable=self.dur_var, command=self.refresh_display)
        self.w["dur_cb"].pack(side=tk.LEFT, padx=2)

        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)

        self.w["font_lbl"] = ttk.Label(toolbar)
        self.w["font_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        fonts = sorted(font.families())
        self.font_combo = ttk.Combobox(toolbar, values=fonts, width=16)
        self.font_combo.set(self.font_family)
        self.font_combo.pack(side=tk.LEFT, padx=2)
        self.font_combo.bind("<<ComboboxSelected>>", self.apply_font)

        self.w["size_lbl"] = ttk.Label(toolbar)
        self.w["size_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        self.size_var = tk.StringVar(value=str(self.font_size))
        self.size_spin = ttk.Spinbox(toolbar, from_=8, to=72, width=4, textvariable=self.size_var)
        self.size_spin.pack(side=tk.LEFT, padx=2)
        self.size_var.trace_add("write", self.apply_font)

        self.w["txt_color_btn"] = ttk.Button(toolbar, command=self.choose_text_color)
        self.w["txt_color_btn"].pack(side=tk.LEFT, padx=2)

        self.w["bg_color_btn"] = ttk.Button(toolbar, command=self.choose_bg_color)
        self.w["bg_color_btn"].pack(side=tk.LEFT, padx=2)

        self.w["margin_lbl"] = ttk.Label(toolbar)
        self.w["margin_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        self.margin_var = tk.StringVar(value=str(self.margin))
        self.margin_spin = ttk.Spinbox(toolbar, from_=0, to=200, width=4,
                                       textvariable=self.margin_var,
                                       command=self.apply_margin)
        self.margin_spin.pack(side=tk.LEFT, padx=2)
        self.margin_var.trace_add("write", self.apply_margin)

        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)
        self.w["status_lbl"] = ttk.Label(toolbar)
        self.w["status_lbl"].pack(side=tk.LEFT, padx=5)

        prog_frame = ttk.Frame(self.root)
        prog_frame.pack(fill=tk.X, padx=5, pady=(0, 2))
        self.progress_bar = ttk.Progressbar(prog_frame, mode="determinate", value=0)
        self.progress_bar.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.progress_bar.bind("<Button-1>", self._on_progress_click)
        self.progress_label = ttk.Label(prog_frame, text="0%", anchor=tk.CENTER,
                                        font=(self.font_family, 10, "bold"))
        self.progress_label.pack(side=tk.RIGHT, padx=(5, 0))
        self.progress_label.bind("<Button-1>", self._on_progress_label_click)

        status_frame = ttk.Frame(self.root, padding=(8, 3))
        status_frame.pack(fill=tk.X, side=tk.BOTTOM)
        self.status_bar = ttk.Label(status_frame, text="", anchor=tk.W)
        self.status_bar.pack(side=tk.LEFT, fill=tk.X, expand=True)

        text_frame = ttk.Frame(self.root, padding=5)
        text_frame.pack(fill=tk.BOTH, expand=True)

        self.text_area = tk.Text(text_frame, wrap=tk.WORD, font=(self.font_family, self.font_size),
                                 fg=self.text_color, bg=self.bg_color, relief=tk.SUNKEN, borderwidth=2)
        self.scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self._on_scrollbar)
        self.text_area.configure(yscrollcommand=self._on_text_scroll)
        self.text_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.text_area.config(state=tk.DISABLED)
        self.text_area.bind("<MouseWheel>", self._on_mousewheel)
        self.text_area.bind("<ButtonRelease-1>", self.update_progress)
        self.text_area.bind("<KeyRelease>", self.update_progress)

    def set_language(self, lang):
        self.current_lang = lang
        self.apply_language()

    # ---- file ----
    def open_file(self, filepath=None, restore_pos=False):
        if filepath is None:
            path = filedialog.askopenfilename(
                title=self._("menu_open"),
                filetypes=[(self._("file_filter"), "*.srt *.sub *.ass *.ssa"), (self._("all_files"), "*.*")]
            )
        else:
            path = filepath
        if not path:
            return
        self.save_position()
        try:
            self.items = SubtitleParser.parse_file(path)
            self.current_filepath = path
            self.root.title(f"{self._('app_title')} - {os.path.basename(path)}")
            self.w["status_lbl"].config(
                text=f"  {os.path.basename(path)}  |  {self._('sub_count').format(len(self.items))}")
            self.refresh_display()
            self.refresh_status_bar()
            if restore_pos:
                self.restore_position()
            self.update_recent_files(path)
            self.update_progress()
            self.rebuild_file_menu()
        except Exception as e:
            self.text_area.config(state=tk.NORMAL)
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, self._("open_err").format(e))
            self.text_area.config(state=tk.DISABLED)
            self.refresh_status_bar()
            self.rebuild_file_menu()

    def on_drop_files(self, paths):
        target = next((p for p in paths
                       if os.path.splitext(p)[1].lower() in SUBTITLE_EXTS), None)
        if target is None and paths:
            target = paths[0]
        if target:
            self.open_file(filepath=target)

    def register_drop_targets(self, widget):
        if isinstance(widget, tk.Menu):
            return
        try:
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", self.on_dnd_drop)
        except Exception:
            return
        for child in widget.winfo_children():
            self.register_drop_targets(child)

    def on_dnd_drop(self, event):
        try:
            paths = self.root.tk.splitlist(event.data)
        except Exception:
            paths = event.data
        if isinstance(paths, str):
            paths = [paths]
        else:
            paths = list(paths)
        self.on_drop_files(paths)
        return "copy"

    def refresh_display(self):
        self.show_time = self.time_var.get()
        self.show_index = self.index_var.get()
        self.show_duration = self.dur_var.get()
        self.text_area.config(state=tk.NORMAL)
        self.text_area.delete(1.0, tk.END)
        if not self.items:
            self.text_area.insert(tk.END, self._("no_file"))
        else:
            for item in self.items:
                parts = []
                if self.show_index:
                    parts.append(f"[{item.index}]")
                if self.show_time:
                    parts.append(f"{item.start} -> {item.end}")
                line = item.text
                if self.show_duration and item.duration:
                    parts.append(f"(dur: {item.duration})")
                prefix = ' '.join(parts)
                if prefix:
                    self.text_area.insert(tk.END, prefix + '\n')
                self.text_area.insert(tk.END, line + '\n\n')
        self.text_area.config(state=tk.DISABLED)
        self.progress_bar.configure(maximum=100)
        self.progress_bar["value"] = 0

    def _format_duration(self, seconds):
        seconds = max(0, int(seconds))
        h, rem = divmod(seconds, 3600)
        m, s = divmod(rem, 60)
        if h:
            return f"{h}:{m:02d}:{s:02d}"
        return f"{m}:{s:02d}"

    def _total_duration(self):
        total = 0.0
        for item in self.items:
            if item.start and item.end:
                try:
                    def to_sec(t):
                        if ":" not in t:
                            return 0.0
                        parts = t.replace(",", ".").split(":")
                        return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
                    total += max(0.0, to_sec(item.end) - to_sec(item.start))
                except Exception:
                    continue
        return total

    def refresh_status_bar(self):
        if not self.items or not self.current_filepath:
            self.status_bar.config(text=self._("status_no_file"))
            return
        name = os.path.basename(self.current_filepath)
        text = f"{self._('status_items').format(len(self.items))}"
        dur = self._total_duration()
        if dur > 0:
            text += f"  |  {self._('status_duration').format(self._format_duration(dur))}"
        self.status_bar.config(text=f"{name}  -  {text}")

    def apply_font(self, *args):
        try:
            self.font_size = int(self.size_var.get())
        except:
            self.font_size = 22
        self.font_family = self.font_combo.get() or "Segoe UI"
        self.set_display_style()

    def apply_margin(self, *args):
        try:
            m = int(self.margin_var.get())
        except Exception:
            m = self.margin
        m = max(0, min(200, m))
        self.margin = m
        self.text_area.config(padx=m)

    def choose_text_color(self):
        color = colorchooser.askcolor(initialcolor=self.text_color, title=self._("text_color_title"))
        if color[1]:
            self.text_color = color[1]
            self.text_area.config(fg=self.text_color)

    def choose_bg_color(self):
        color = colorchooser.askcolor(initialcolor=self.bg_color, title=self._("bg_color_title"))
        if color[1]:
            self.bg_color = color[1]
            self.text_area.config(bg=self.bg_color)

    # ---- settings persistence ----
    def load_settings(self):
        try:
            cfg = self.load_config()
        except Exception:
            cfg = {}
        settings = cfg.get("settings", {})
        size = settings.get("font_size", self.font_size)
        try:
            size = int(size)
        except Exception:
            size = self.font_size
        family = settings.get("font_family") or self.font_family
        if family not in font.families():
            family = self.font_family
        self.font_size = size
        self.font_family = family
        self.text_color = settings.get("text_color", self.text_color)
        self.bg_color = settings.get("bg_color", self.bg_color)
        try:
            self.margin = max(0, min(200, int(settings.get("margin", self.margin))))
        except Exception:
            pass
        lang = settings.get("language")
        if lang in ("zh", "en"):
            self.current_lang = lang
        self.font_combo.set(family)
        self.size_var.set(str(size))
        self.margin_var.set(str(self.margin))
        self.set_display_style()
        if lang in ("zh", "en"):
            self.apply_language()

    def set_display_style(self):
        self.text_area.config(font=(self.font_family, self.font_size),
                              fg=self.text_color, bg=self.bg_color,
                              padx=self.margin)
        self.apply_ui_font()

    def apply_ui_font(self):
        fam = self.font_family
        sz = self.ui_font_size
        try:
            self.menubar.config(font=(fam, sz))
            for m in (self.file_menu, self.display_menu, self.lang_menu, self.help_menu):
                m.config(font=(fam, sz))
        except Exception:
            pass
        try:
            style = ttk.Style(self.root)
            for cls in ("TLabel", "TButton", "TCheckbutton", "TSpinbox",
                        "TCombobox", "TEntry", "TProgressbar"):
                style.configure(cls, font=(fam, sz))
        except Exception:
            pass
        try:
            self.progress_label.config(font=(fam, sz, "bold"))
        except Exception:
            pass

    def save_settings(self):
        try:
            cfg = self.load_config()
        except Exception:
            cfg = {}
        cfg["settings"] = {
            "font_family": self.font_family,
            "font_size": self.font_size,
            "text_color": self.text_color,
            "bg_color": self.bg_color,
            "margin": self.margin,
            "language": self.current_lang,
        }
        self.save_config(cfg)

    def on_close(self):
        self.save_position()
        self.save_settings()
        self.root.destroy()

    # ---- export ----
    def build_export_text(self):
        lines = []
        if not self.items:
            return lines
        for item in self.items:
            parts = []
            if self.index_var.get():
                parts.append(f"[{item.index}]")
            if self.time_var.get():
                parts.append(f"{item.start} -> {item.end}")
            if self.dur_var.get() and item.duration:
                parts.append(f"(dur: {item.duration})")
            prefix = ' '.join(parts)
            if prefix:
                lines.append(prefix)
            lines.append(item.text)
            lines.append('')
        return lines

    def export_content(self):
        if not self.items:
            messagebox.showinfo(self._("export_title"), self._("export_none"))
            return
        base = os.path.splitext(os.path.basename(self.current_filepath))[0]
        path = filedialog.asksaveasfilename(
            initialfile=base,
            defaultextension=".txt",
            filetypes=[
                (self._("filetype_txt"), "*.txt"),
                (self._("filetype_docx"), "*.docx"),
                (self._("all_files"), "*.*"),
            ],
            title=self._("export_title")
        )
        if not path:
            return
        ext = os.path.splitext(path)[1].lower()
        try:
            text_lines = self.build_export_text()
            content = '\n'.join(text_lines)
            if ext == ".txt":
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(content)
            elif ext == ".docx":
                if not HAS_DOCX:
                    messagebox.showerror(self._("export_err"),
                        "python-docx is not installed.\nRun: pip install python-docx")
                    return
                doc = Document()
                style = doc.styles['Normal']
                style.font.size = Pt(self.font_size)
                for line in text_lines:
                    doc.add_paragraph(line)
                doc.save(path)
            else:
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(content)
            messagebox.showinfo(self._("export_title"), self._("export_ok").format(path))
        except Exception as e:
            messagebox.showerror(self._("export_err"), str(e))

    # ---- config ----
    def get_config_dir(self):
        if sys.platform == "win32":
            base = os.environ.get("APPDATA", os.path.expanduser("~"))
        else:
            base = os.path.expanduser("~/.config")
        cfg_dir = os.path.join(base, "SubtitleReader")
        os.makedirs(cfg_dir, exist_ok=True)
        return cfg_dir

    def get_config_path(self):
        return os.path.join(self.get_config_dir(), "config.json")

    def load_config(self):
        try:
            with open(self.get_config_path(), encoding="utf-8") as f:
                return json.load(f)
        except:
            return {"recent_files": [], "positions": {}}

    def save_config(self, cfg):
        try:
            with open(self.get_config_path(), "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False, indent=2)
        except:
            pass

    def update_recent_files(self, filepath):
        cfg = self.load_config()
        recent = cfg.get("recent_files", [])
        if filepath in recent:
            recent.remove(filepath)
        recent.insert(0, filepath)
        cfg["recent_files"] = recent[:10]
        self.save_config(cfg)

    # ---- position ----
    def get_current_index(self):
        if not self.items:
            return 0
        try:
            frac = self.text_area.yview()[0]
            return max(0, min(int(frac * len(self.items)), len(self.items) - 1))
        except:
            return 0

    def save_position(self):
        if not self.current_filepath or not self.items:
            return
        cfg = self.load_config()
        if "positions" not in cfg:
            cfg["positions"] = {}
        idx = self.get_current_index()
        frac = self.text_area.yview()[0]
        cfg["positions"][self.current_filepath] = {"index": idx, "fraction": frac}
        self.save_config(cfg)

    def restore_position(self):
        if not self.current_filepath:
            return
        cfg = self.load_config()
        pos = cfg.get("positions", {}).get(self.current_filepath)
        if pos:
            frac = pos.get("fraction", 0.0)
            self.root.after_idle(lambda: self.text_area.yview_moveto(frac))

    # ---- close ----
    def close_file(self):
        if not self.items:
            return
        self.save_position()
        self.items = []
        self.current_filepath = None
        self.root.title(self._("app_title"))
        self.w["status_lbl"].config(text="")
        self.refresh_display()
        self.refresh_status_bar()
        self.update_progress()
        self.rebuild_file_menu()

    # ---- recent files ----
    def open_recent_file(self, filepath):
        if not os.path.exists(filepath):
            cfg = self.load_config()
            recent = cfg.get("recent_files", [])
            if filepath in recent:
                recent.remove(filepath)
            cfg["recent_files"] = recent
            self.save_config(cfg)
            self.rebuild_file_menu()
            return
        self.open_file(filepath=filepath, restore_pos=True)

    # ---- progress ----
    def update_progress(self, *args):
        if not self.items:
            self.progress_bar.configure(maximum=100)
            self.progress_bar["value"] = 0
            self.progress_label.config(text="0%")
            return
        try:
            if self.display_mode.get() == "percent":
                frac = self.text_area.yview()[0]
                val = int(frac * 100)
                self.progress_bar.configure(maximum=100)
                self.progress_bar["value"] = val
                self.progress_label.config(text=f"{val}%")
            else:
                yv = self.text_area.yview()
                visible = yv[1] - yv[0]
                if visible <= 0:
                    visible = 0.01
                total = max(1, int(1.0 / visible + 0.5))
                cur = int(yv[0] * total) + 1
                cur = max(1, min(cur, total))
                self.progress_bar.configure(maximum=total)
                self.progress_bar["value"] = cur
                self.progress_label.config(text=self._("page_fmt").format(cur, total))
        except:
            self.progress_bar.configure(maximum=100)
            self.progress_bar["value"] = 0
            self.progress_label.config(text="0%")

    def _on_scrollbar(self, *args):
        self.text_area.yview(*args)
        self.update_progress()

    def _on_text_scroll(self, *args):
        self.scrollbar.set(*args)
        self.update_progress()

    def _on_mousewheel(self, event):
        self.root.after(10, self.update_progress)

    # ---- menu rebuild ----
    def rebuild_file_menu(self):
        self.file_menu.delete(0, "end")
        has_file = bool(self.items)

        self.file_menu.add_command(label=self._("menu_open"), command=self.open_file, accelerator="Ctrl+O")
        self.file_menu.add_command(
            label=self._("menu_close"), command=self.close_file,
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_separator()
        self.file_menu.add_command(
            label=self._("menu_export"), command=self.export_content, accelerator="Ctrl+E",
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_separator()

        cfg = self.load_config()
        recent = cfg.get("recent_files", [])
        positions = cfg.get("positions", {})
        self._last_rebuild_recent = list(recent)
        if recent:
            self.file_menu.add_command(
                label="── " + self._("menu_recent") + " ──",
                state=tk.DISABLED
            )
            for fpath in recent:
                exists = os.path.exists(fpath)
                label = os.path.basename(fpath)
                pos = positions.get(fpath, {})
                frac = pos.get("fraction", 0.0)
                pct = min(100, max(0, int(frac * 100)))
                if 0 < pct < 100:
                    label = f"[{pct}%] {label}"
                elif pct >= 100:
                    label = f"[100%] {label}"
                if not exists:
                    label += " " + self._("file_missing")
                if exists:
                    self.file_menu.add_command(
                        label=label,
                        command=lambda p=fpath: self.open_recent_file(p)
                    )
                else:
                    self.file_menu.add_command(label=label, state=tk.DISABLED)
            self.file_menu.add_separator()
            self.file_menu.add_command(
                label=self._("clear_recent"),
                command=self.clear_recent_files
            )

        self.file_menu.add_separator()
        self.file_menu.add_command(label=self._("menu_exit"), command=self.root.quit)

    def clear_recent_files(self):
        if not messagebox.askyesno(self._("clear_recent"), self._("clear_recent_confirm")):
            return
        cfg = self.load_config()
        cfg["recent_files"] = []
        self.save_config(cfg)
        self.rebuild_file_menu()

    def _on_progress_click(self, event):
        if not self.items:
            return
        width = self.progress_bar.winfo_width()
        if width <= 0:
            return
        fraction = event.x / width
        if self.display_mode.get() == "percent":
            target_idx = int(fraction * len(self.items))
            target_idx = max(0, min(target_idx, len(self.items) - 1))
            target_frac = target_idx / len(self.items)
            self.text_area.yview_moveto(target_frac)
        else:
            yv = self.text_area.yview()
            visible = yv[1] - yv[0]
            if visible <= 0:
                visible = 0.01
            total = max(1, int(1.0 / visible + 0.5))
            target_page = int(fraction * total) + 1
            target_page = max(1, min(target_page, total))
            target_frac = (target_page - 1) / total
            self.text_area.yview_moveto(target_frac)
        self.update_progress()

    def _on_progress_label_click(self, event):
        if not self.items:
            return
        if self.display_mode.get() == "percent":
            cur_val = self.progress_bar["value"]
            result = simpledialog.askinteger(
                self._("jump_to"), self._("jump_to_hint"),
                minvalue=0, maxvalue=100,
                initialvalue=int(cur_val),
                parent=self.root
            )
            if result is not None:
                fraction = result / 100.0
                target_idx = int(fraction * len(self.items))
                target_idx = max(0, min(target_idx, len(self.items) - 1))
                target_frac = target_idx / len(self.items)
                self.text_area.yview_moveto(target_frac)
        else:
            yv = self.text_area.yview()
            visible = yv[1] - yv[0]
            if visible <= 0:
                visible = 0.01
            total = max(1, int(1.0 / visible + 0.5))
            cur_page = int(yv[0] * total) + 1
            cur_page = max(1, min(cur_page, total))
            result = simpledialog.askinteger(
                self._("jump_to"), self._("page_jump_hint").format(total),
                minvalue=1, maxvalue=total,
                initialvalue=cur_page,
                parent=self.root
            )
            if result is not None:
                target_frac = (result - 1) / total
                self.text_area.yview_moveto(target_frac)
        self.update_progress()

    def on_display_mode_change(self):
        self.update_progress()

    def show_about(self):
        win = tk.Toplevel(self.root)
        win.title(self._("about_title"))
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        lines = self._("about_info").split('\n')
        for i, line in enumerate(lines):
            lbl = tk.Label(win, text=line, font=(self.font_family, self.ui_font_size), anchor=tk.CENTER)
            lbl.pack(fill=tk.X, padx=30, pady=(10 if i == 0 else 0, 0))
        tk.Button(win, text=self._("menu_close"), command=win.destroy,
                  width=10, font=(self.font_family, self.ui_font_size)).pack(pady=15)
        win.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() - win.winfo_width()) // 2
        y = self.root.winfo_rooty() + 50
        win.geometry(f"+{x}+{y}")

    def set_window_icon(self):
        base = getattr(sys, '_MEIPASS', os.path.dirname(__file__))
        ico_path = os.path.join(base, "icon.ico")
        png_path = os.path.join(base, "icon.png")
        if os.path.exists(ico_path):
            try:
                self.root.iconbitmap(ico_path)
            except:
                pass
        elif os.path.exists(png_path):
            try:
                img = tk.PhotoImage(file=png_path)
                self.root.iconphoto(True, img)
            except:
                pass

if __name__ == "__main__":
    enable_high_dpi()
    dnd_root = None
    if HAS_DND:
        try:
            dnd_root = TkinterDnD.Tk()
        except Exception:
            HAS_DND = False
            dnd_root = None
    root = dnd_root if dnd_root is not None else tk.Tk()
    scale = get_dpi_scale()
    try:
        root.tk.call("tk", "scaling", scale if scale != 1.0 else 1.0)
    except Exception:
        pass
    work = get_work_area()
    width = int(1200 * scale) if scale != 1.0 else 1200
    if work and work["height"] > 0:
        max_w = work["width"]
        if width > max_w:
            width = max_w
        x = work["x"] + (max_w - width) // 2
        root.withdraw()
        root.geometry(f"{width}x{work['height']}+{x}+{work['y']}")
        root.update_idletasks()
        root.deiconify()
        root.update()
    else:
        if scale != 1.0:
            root.geometry(f"{int(1200 * scale)}x{int(1200 * scale)}")
        else:
            root.geometry("1200x1200")
    app = SubtitleReader(root)
    if work and work["height"] > 0:
        root.update_idletasks()
        try:
            rect = _window_outer(root)
            if rect:
                overflow = (rect[1] + rect[3]) - (work["y"] + work["height"])
                if overflow > 0:
                    height = max(200, work["height"] - overflow)
                    root.geometry(f"{width}x{height}+{x}+{work['y']}")
                    root.update_idletasks()
        except Exception:
            pass
    root.mainloop()
