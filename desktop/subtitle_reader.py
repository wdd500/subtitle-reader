import sys
import ctypes
import tkinter as tk
from tkinter import filedialog, colorchooser, font, ttk, messagebox
import re
import os
import json

try:
    from docx import Document
    from docx.shared import Pt
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


LANG = {
    "zh": {
        "app_title": "字幕阅读器",
        "menu_file": "文件",
        "menu_open": "打开...",
        "menu_export": "导出...",
        "menu_exit": "退出",
        "menu_lang": "语言",
        "lang_zh": "中文",
        "lang_en": "English",
        "open_file": "打开文件",
        "timecode": "时间码",
        "line_no": "行号",
        "duration": "时长",
        "font": "字体：",
        "size": "字号：",
        "text_color": "文字颜色",
        "bg_color": "背景颜色",
        "no_file": "未加载字幕。\n打开 .srt/.sub/.ass 文件开始阅读。",
        "export_title": "导出字幕内容",
        "export_ok": "已导出到：\n{}",
        "export_err": "导出错误",
        "export_none": "未加载字幕文件。",
        "text_color_title": "文字颜色",
        "bg_color_title": "背景颜色",
        "sub_count": "{} 条字幕",
        "open_err": "打开文件出错：{}",
        "file_filter": "字幕文件",
        "all_files": "所有文件",
        "filetype_txt": "文本文件",
        "filetype_docx": "Word 文档",
        "export_fmt": "导出格式",
        "export_txt": "文本 (.txt)",
        "export_docx": "Word (.docx)",
        "menu_close": "关闭",
        "menu_recent": "最近打开",
        "file_missing": "(文件不存在)",
    },
    "en": {
        "app_title": "Subtitle Reader",
        "menu_file": "File",
        "menu_open": "Open...",
        "menu_export": "Export...",
        "menu_exit": "Exit",
        "menu_lang": "Language",
        "lang_zh": "中文",
        "lang_en": "English",
        "open_file": "Open File",
        "timecode": "Timecode",
        "line_no": "Line No.",
        "duration": "Duration",
        "font": "Font:",
        "size": "Size:",
        "text_color": "Text Color",
        "bg_color": "Bg Color",
        "no_file": "No subtitles loaded.\nOpen a .srt / .sub / .ass file to begin.",
        "export_title": "Export Subtitle Content",
        "export_ok": "Exported to:\n{}",
        "export_err": "Export Error",
        "export_none": "No subtitle file loaded.",
        "text_color_title": "Text Color",
        "bg_color_title": "Background Color",
        "sub_count": "{} subtitles",
        "open_err": "Error opening file: {}",
        "file_filter": "Subtitle files",
        "all_files": "All files",
        "filetype_txt": "Text file",
        "filetype_docx": "Word document",
        "export_fmt": "Export Format",
        "export_txt": "Text (.txt)",
        "export_docx": "Word (.docx)",
        "menu_close": "Close",
        "menu_recent": "Recent Files",
        "file_missing": "(file missing)",
    },
}


def enable_high_dpi():
    if sys.platform != "win32":
        return
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


def get_dpi_scale():
    try:
        hwnd = ctypes.windll.user32.GetDesktopWindow()
        dpi = ctypes.windll.user32.GetDpiForWindow(hwnd)
        return dpi / 96.0
    except Exception:
        return 1.0


class SubtitleItem:
    def __init__(self, index=0, start="", end="", text="", duration=""):
        self.index = index
        self.start = start
        self.end = end
        self.text = text
        self.duration = duration


class SubtitleParser:
    @staticmethod
    def parse_srt(content):
        items = []
        blocks = re.split(r'\n\s*\n', content.strip())
        for block in blocks:
            lines = block.strip().split('\n')
            if len(lines) >= 3:
                try:
                    idx = int(lines[0])
                    time_match = re.match(
                        r'(\d+:\d+:\d+[,.]\d+)\s*-->\s*(\d+:\d+:\d+[,.]\d+)', lines[1])
                    if time_match:
                        start = time_match.group(1).replace(',', '.')
                        end = time_match.group(2).replace(',', '.')
                        text = '\n'.join(lines[2:])
                        duration = SubtitleParser._calc_duration(start, end)
                        items.append(SubtitleItem(idx, start, end, text, duration))
                except ValueError:
                    continue
        return items

    @staticmethod
    def parse_sub(content):
        items = []
        lines = content.strip().split('\n')
        for i, line in enumerate(lines):
            line = line.strip()
            if not line:
                continue
            match = re.match(r'\{(\d+)\}\{(\d+)\}(.*)', line)
            if match:
                start_frame = match.group(1)
                end_frame = match.group(2)
                text = match.group(3).replace('|', '\n')
                duration = str(int(end_frame) - int(start_frame)) + ' frames'
                items.append(SubtitleItem(i + 1, start_frame, end_frame, text, duration))
        return items

    @staticmethod
    def parse_ass(content):
        items = []
        events_section = False
        format_line = None
        for line in content.split('\n'):
            line = line.strip()
            if line.startswith('[Events]'):
                events_section = True
                continue
            if events_section:
                if line.startswith('Format:'):
                    format_line = line[7:].strip().split(',')
                    format_line = [f.strip() for f in format_line]
                    continue
                if line.startswith('Dialogue:'):
                    parts = line[9:].strip()
                    if format_line:
                        vals = SubtitleParser._split_ass(parts)
                        data = {}
                        for j, key in enumerate(format_line):
                            if j < len(vals):
                                data[key] = vals[j]
                        start = data.get('Start', '')
                        end = data.get('End', '')
                        text = data.get('Text', '')
                        text = re.sub(r'\{[^}]*\}', '', text)
                        text = text.replace('\\N', '\n').replace('\\n', '\n')
                        duration = SubtitleParser._calc_duration_ass(start, end)
                        items.append(SubtitleItem(len(items) + 1, start, end, text, duration))
        return items

    @staticmethod
    def _split_ass(line):
        parts = []
        current = ''
        in_braces = 0
        for ch in line:
            if ch == '{':
                in_braces += 1
                current += ch
            elif ch == '}':
                in_braces -= 1
                current += ch
            elif ch == ',' and in_braces == 0:
                parts.append(current.strip())
                current = ''
            else:
                current += ch
        parts.append(current.strip())
        return parts

    @staticmethod
    def _calc_duration(start, end):
        def to_seconds(t):
            parts = t.replace(',', '.').split(':')
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
        try:
            d = to_seconds(end) - to_seconds(start)
            return f'{d:.2f}s'
        except:
            return ''

    @staticmethod
    def _calc_duration_ass(start, end):
        def to_seconds(t):
            parts = t.replace(',', '.').split(':')
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
        try:
            d = to_seconds(end) - to_seconds(start)
            return f'{d:.2f}s'
        except:
            return ''

    @staticmethod
    def parse_file(filepath):
        ext = os.path.splitext(filepath)[1].lower()
        with open(filepath, encoding='utf-8') as f:
            content = f.read()
        if ext == '.srt':
            return SubtitleParser.parse_srt(content)
        elif ext == '.sub':
            return SubtitleParser.parse_sub(content)
        elif ext in ('.ass', '.ssa'):
            return SubtitleParser.parse_ass(content)
        else:
            return []


class SubtitleReader:
    def __init__(self, root):
        self.root = root
        self.current_lang = "zh"
        self.items = []
        self.current_filepath = None
        self.show_time = False
        self.show_index = False
        self.show_duration = False
        self.text_color = "#000000"
        self.bg_color = "#FFFFFF"
        self.font_size = 22
        self.font_family = "Segoe UI"
        self.w = {}
        self._last_rebuild_recent = []
        self.build_ui()
        self.apply_language()

    # ---- i18n helpers ----
    def _(self, key):
        return LANG[self.current_lang].get(key, key)

    def apply_language(self):
        self.root.title(self._("app_title"))
        self.rebuild_file_menu()

        self.lang_menu.delete(0, "end")
        self.lang_menu.add_command(label=self._("lang_zh"), command=lambda: self.set_language("zh"))
        self.lang_menu.add_command(label=self._("lang_en"), command=lambda: self.set_language("en"))

        self.menubar.delete(0, "end")
        self.menubar.add_cascade(label=self._("menu_file"), menu=self.file_menu)
        self.menubar.add_cascade(label=self._("menu_lang"), menu=self.lang_menu)
        self.root.config(menu=self.menubar)

        self.w["open_btn"].config(text=self._("open_file"))
        self.w["time_cb"].config(text=self._("timecode"))
        self.w["idx_cb"].config(text=self._("line_no"))
        self.w["dur_cb"].config(text=self._("duration"))
        self.w["font_lbl"].config(text=self._("font"))
        self.w["size_lbl"].config(text=self._("size"))
        self.w["txt_color_btn"].config(text=self._("text_color"))
        self.w["bg_color_btn"].config(text=self._("bg_color"))

        self.refresh_display()

    # ---- build UI ----
    def build_ui(self):
        self.menubar = tk.Menu(self.root)
        self.file_menu = tk.Menu(self.menubar, tearoff=0)
        self.file_menu.add_command(label="", command=self.open_file, accelerator="Ctrl+O")
        self.file_menu.add_separator()
        self.file_menu.add_command(label="", command=self.export_content, accelerator="Ctrl+E")
        self.file_menu.add_separator()
        self.file_menu.add_command(label="", command=self.root.quit)
        self.menubar.add_cascade(label="", menu=self.file_menu)

        self.lang_menu = tk.Menu(self.menubar, tearoff=0)
        self.lang_menu.add_command(label="", command=lambda: self.set_language("zh"))
        self.lang_menu.add_command(label="", command=lambda: self.set_language("en"))
        self.menubar.add_cascade(label="", menu=self.lang_menu)

        self.root.config(menu=self.menubar)
        self.set_window_icon()
        self.root.bind_all("<Control-o>", lambda e: self.open_file())
        self.root.bind_all("<Control-e>", lambda e: self.export_content())

        toolbar = ttk.Frame(self.root)
        toolbar.pack(fill=tk.X, padx=3, pady=3)

        self.w["open_btn"] = ttk.Button(toolbar, command=self.open_file)
        self.w["open_btn"].pack(side=tk.LEFT, padx=2)

        self.time_var = tk.BooleanVar()
        self.w["time_cb"] = ttk.Checkbutton(toolbar, variable=self.time_var, command=self.refresh_display)
        self.w["time_cb"].pack(side=tk.LEFT, padx=2)

        self.index_var = tk.BooleanVar()
        self.w["idx_cb"] = ttk.Checkbutton(toolbar, variable=self.index_var, command=self.refresh_display)
        self.w["idx_cb"].pack(side=tk.LEFT, padx=2)

        self.dur_var = tk.BooleanVar()
        self.w["dur_cb"] = ttk.Checkbutton(toolbar, variable=self.dur_var, command=self.refresh_display)
        self.w["dur_cb"].pack(side=tk.LEFT, padx=2)

        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)

        self.w["font_lbl"] = ttk.Label(toolbar)
        self.w["font_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        fonts = sorted(font.families())
        self.font_combo = ttk.Combobox(toolbar, values=fonts, width=16)
        self.font_combo.set(self.font_family)
        self.font_combo.pack(side=tk.LEFT, padx=2)
        self.font_combo.bind("<<ComboboxSelected>>", self.apply_font)

        self.w["size_lbl"] = ttk.Label(toolbar)
        self.w["size_lbl"].pack(side=tk.LEFT, padx=(5, 0))
        self.size_var = tk.StringVar(value=str(self.font_size))
        self.size_spin = ttk.Spinbox(toolbar, from_=8, to=72, width=4, textvariable=self.size_var)
        self.size_spin.pack(side=tk.LEFT, padx=2)
        self.size_var.trace_add("write", self.apply_font)

        self.w["txt_color_btn"] = ttk.Button(toolbar, command=self.choose_text_color)
        self.w["txt_color_btn"].pack(side=tk.LEFT, padx=2)

        self.w["bg_color_btn"] = ttk.Button(toolbar, command=self.choose_bg_color)
        self.w["bg_color_btn"].pack(side=tk.LEFT, padx=2)

        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)
        self.w["status_lbl"] = ttk.Label(toolbar)
        self.w["status_lbl"].pack(side=tk.LEFT, padx=5)

        prog_frame = ttk.Frame(self.root)
        prog_frame.pack(fill=tk.X, padx=5, pady=(0, 2))
        self.progress_bar = ttk.Progressbar(prog_frame, mode="determinate", value=0)
        self.progress_bar.pack(fill=tk.X, expand=True)

        text_frame = ttk.Frame(self.root, padding=5)
        text_frame.pack(fill=tk.BOTH, expand=True)

        self.text_area = tk.Text(text_frame, wrap=tk.WORD, font=(self.font_family, self.font_size),
                                 fg=self.text_color, bg=self.bg_color, relief=tk.SUNKEN, borderwidth=2)
        self.scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self._on_scrollbar)
        self.text_area.configure(yscrollcommand=self._on_text_scroll)
        self.text_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.text_area.config(state=tk.DISABLED)
        self.text_area.bind("<MouseWheel>", self._on_mousewheel)
        self.text_area.bind("<ButtonRelease-1>", self.update_progress)
        self.text_area.bind("<KeyRelease>", self.update_progress)

    def set_language(self, lang):
        self.current_lang = lang
        self.apply_language()

    # ---- file ----
    def open_file(self, filepath=None, restore_pos=False):
        if filepath is None:
            path = filedialog.askopenfilename(
                title=self._("menu_open"),
                filetypes=[(self._("file_filter"), "*.srt *.sub *.ass *.ssa"), (self._("all_files"), "*.*")]
            )
        else:
            path = filepath
        if not path:
            return
        self.save_position()
        try:
            self.items = SubtitleParser.parse_file(path)
            self.current_filepath = path
            self.root.title(f"{self._('app_title')} - {os.path.basename(path)}")
            self.w["status_lbl"].config(
                text=f"  {os.path.basename(path)}  |  {self._('sub_count').format(len(self.items))}")
            self.refresh_display()
            if restore_pos:
                self.restore_position()
            self.update_recent_files(path)
            self.update_progress()
            self.rebuild_file_menu()
        except Exception as e:
            self.text_area.config(state=tk.NORMAL)
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, self._("open_err").format(e))
            self.text_area.config(state=tk.DISABLED)
            self.rebuild_file_menu()

    def refresh_display(self):
        self.show_time = self.time_var.get()
        self.show_index = self.index_var.get()
        self.show_duration = self.dur_var.get()
        self.text_area.config(state=tk.NORMAL)
        self.text_area.delete(1.0, tk.END)
        if not self.items:
            self.text_area.insert(tk.END, self._("no_file"))
        else:
            for item in self.items:
                parts = []
                if self.show_index:
                    parts.append(f"[{item.index}]")
                if self.show_time:
                    parts.append(f"{item.start} -> {item.end}")
                line = item.text
                if self.show_duration and item.duration:
                    parts.append(f"(dur: {item.duration})")
                prefix = ' '.join(parts)
                if prefix:
                    self.text_area.insert(tk.END, prefix + '\n')
                self.text_area.insert(tk.END, line + '\n\n')
        self.text_area.config(state=tk.DISABLED)
        self.progress_bar["value"] = 0

    def apply_font(self, *args):
        try:
            self.font_size = int(self.size_var.get())
        except:
            self.font_size = 22
        self.font_family = self.font_combo.get() or "Segoe UI"
        self.text_area.config(font=(self.font_family, self.font_size))

    def choose_text_color(self):
        color = colorchooser.askcolor(initialcolor=self.text_color, title=self._("text_color_title"))
        if color[1]:
            self.text_color = color[1]
            self.text_area.config(fg=self.text_color)

    def choose_bg_color(self):
        color = colorchooser.askcolor(initialcolor=self.bg_color, title=self._("bg_color_title"))
        if color[1]:
            self.bg_color = color[1]
            self.text_area.config(bg=self.bg_color)

    # ---- export ----
    def build_export_text(self):
        lines = []
        if not self.items:
            return lines
        for item in self.items:
            parts = []
            if self.index_var.get():
                parts.append(f"[{item.index}]")
            if self.time_var.get():
                parts.append(f"{item.start} -> {item.end}")
            if self.dur_var.get() and item.duration:
                parts.append(f"(dur: {item.duration})")
            prefix = ' '.join(parts)
            if prefix:
                lines.append(prefix)
            lines.append(item.text)
            lines.append('')
        return lines

    def export_content(self):
        if not self.items:
            messagebox.showinfo(self._("export_title"), self._("export_none"))
            return
        base = os.path.splitext(os.path.basename(self.current_filepath))[0]
        path = filedialog.asksaveasfilename(
            initialfile=base,
            defaultextension=".txt",
            filetypes=[
                (self._("filetype_txt"), "*.txt"),
                (self._("filetype_docx"), "*.docx"),
                (self._("all_files"), "*.*"),
            ],
            title=self._("export_title")
        )
        if not path:
            return
        ext = os.path.splitext(path)[1].lower()
        try:
            text_lines = self.build_export_text()
            content = '\n'.join(text_lines)
            if ext == ".txt":
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(content)
            elif ext == ".docx":
                if not HAS_DOCX:
                    messagebox.showerror(self._("export_err"),
                        "python-docx is not installed.\nRun: pip install python-docx")
                    return
                doc = Document()
                style = doc.styles['Normal']
                style.font.size = Pt(self.font_size)
                for line in text_lines:
                    doc.add_paragraph(line)
                doc.save(path)
            else:
                with open(path, 'w', encoding='utf-8') as f:
                    f.write(content)
            messagebox.showinfo(self._("export_title"), self._("export_ok").format(path))
        except Exception as e:
            messagebox.showerror(self._("export_err"), str(e))

    # ---- config ----
    def get_config_dir(self):
        if sys.platform == "win32":
            base = os.environ.get("APPDATA", os.path.expanduser("~"))
        else:
            base = os.path.expanduser("~/.config")
        cfg_dir = os.path.join(base, "SubtitleReader")
        os.makedirs(cfg_dir, exist_ok=True)
        return cfg_dir

    def get_config_path(self):
        return os.path.join(self.get_config_dir(), "config.json")

    def load_config(self):
        try:
            with open(self.get_config_path(), encoding="utf-8") as f:
                return json.load(f)
        except:
            return {"recent_files": [], "positions": {}}

    def save_config(self, cfg):
        try:
            with open(self.get_config_path(), "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False, indent=2)
        except:
            pass

    def update_recent_files(self, filepath):
        cfg = self.load_config()
        recent = cfg.get("recent_files", [])
        if filepath in recent:
            recent.remove(filepath)
        recent.insert(0, filepath)
        cfg["recent_files"] = recent[:10]
        self.save_config(cfg)

    # ---- position ----
    def get_current_index(self):
        if not self.items:
            return 0
        try:
            frac = self.text_area.yview()[0]
            return max(0, min(int(frac * len(self.items)), len(self.items) - 1))
        except:
            return 0

    def save_position(self):
        if not self.current_filepath or not self.items:
            return
        cfg = self.load_config()
        if "positions" not in cfg:
            cfg["positions"] = {}
        idx = self.get_current_index()
        frac = self.text_area.yview()[0]
        cfg["positions"][self.current_filepath] = {"index": idx, "fraction": frac}
        self.save_config(cfg)

    def restore_position(self):
        if not self.current_filepath:
            return
        cfg = self.load_config()
        pos = cfg.get("positions", {}).get(self.current_filepath)
        if pos:
            frac = pos.get("fraction", 0.0)
            self.root.after_idle(lambda: self.text_area.yview_moveto(frac))

    # ---- close ----
    def close_file(self):
        if not self.items:
            return
        self.save_position()
        self.items = []
        self.current_filepath = None
        self.root.title(self._("app_title"))
        self.w["status_lbl"].config(text="")
        self.refresh_display()
        self.update_progress()
        self.rebuild_file_menu()

    # ---- recent files ----
    def open_recent_file(self, filepath):
        if not os.path.exists(filepath):
            cfg = self.load_config()
            recent = cfg.get("recent_files", [])
            if filepath in recent:
                recent.remove(filepath)
            cfg["recent_files"] = recent
            self.save_config(cfg)
            self.rebuild_file_menu()
            return
        self.open_file(filepath=filepath, restore_pos=True)

    # ---- progress ----
    def update_progress(self, *args):
        if not self.items:
            self.progress_bar["value"] = 0
            return
        try:
            frac = self.text_area.yview()[0]
            self.progress_bar["value"] = int(frac * 100)
        except:
            self.progress_bar["value"] = 0

    def _on_scrollbar(self, *args):
        self.text_area.yview(*args)
        self.update_progress()

    def _on_text_scroll(self, *args):
        self.scrollbar.set(*args)
        self.update_progress()

    def _on_mousewheel(self, event):
        self.root.after(10, self.update_progress)

    # ---- menu rebuild ----
    def rebuild_file_menu(self):
        self.file_menu.delete(0, "end")
        has_file = bool(self.items)

        self.file_menu.add_command(label=self._("menu_open"), command=self.open_file, accelerator="Ctrl+O")
        self.file_menu.add_command(
            label=self._("menu_close"), command=self.close_file,
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_separator()
        self.file_menu.add_command(
            label=self._("menu_export"), command=self.export_content, accelerator="Ctrl+E",
            state=tk.NORMAL if has_file else tk.DISABLED
        )
        self.file_menu.add_separator()

        cfg = self.load_config()
        recent = cfg.get("recent_files", [])
        self._last_rebuild_recent = list(recent)
        if recent:
            self.file_menu.add_command(
                label="── " + self._("menu_recent") + " ──",
                state=tk.DISABLED
            )
            for fpath in recent:
                exists = os.path.exists(fpath)
                label = os.path.basename(fpath)
                if not exists:
                    label += " " + self._("file_missing")
                if exists:
                    self.file_menu.add_command(
                        label=label,
                        command=lambda p=fpath: self.open_recent_file(p)
                    )
                else:
                    self.file_menu.add_command(label=label, state=tk.DISABLED)

        self.file_menu.add_separator()
        self.file_menu.add_command(label=self._("menu_exit"), command=self.root.quit)



    def set_window_icon(self):
        ico_path = os.path.join(os.path.dirname(__file__), "icon.ico")
        png_path = os.path.join(os.path.dirname(__file__), "icon.png")
        if os.path.exists(ico_path):
            try:
                self.root.iconbitmap(ico_path)
            except:
                pass
        elif os.path.exists(png_path):
            try:
                img = tk.PhotoImage(file=png_path)
                self.root.iconphoto(True, img)
            except:
                pass

if __name__ == "__main__":
    enable_high_dpi()
    root = tk.Tk()
    scale = get_dpi_scale()
    if scale != 1.0:
        try:
            root.tk.call("tk", "scaling", scale)
        except Exception:
            pass
        root.geometry(f"{int(1100 * scale)}x{int(700 * scale)}")
    else:
        root.geometry("1100x700")
    app = SubtitleReader(root)
    root.mainloop()
